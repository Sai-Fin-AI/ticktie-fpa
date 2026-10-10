"""Profit centres and the product hierarchy: sales by profit centre and SKU (or brand), costs by cost centre."""
import io
import sys
import unittest
from decimal import Decimal as D
from pathlib import Path

from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))

import validate  # noqa: E402
from agents.deck import make_deck, products_headline  # noqa: E402
from checker import short_name  # noqa: E402
from finance import (CC, UnmappedSKU, aggregate, apply_products, breakdown, check_products, load_csv,  # noqa: E402
                     load_product_master, merge_datasets, product_margins, review_lines, sample_csv, sample_pl_csv,
                     sample_product_master_csv)
from gl_panel import product_workbook  # noqa: E402
from ingest import guess_roles, table_csv  # noqa: E402
from pack import lock_pack, pack_excel  # noqa: E402
from record_demo import demo_pack, demo_rows  # noqa: E402
from test_deck import slide_texts  # noqa: E402

APP = ROOT / 'app.py'
HEADER = 'month,country,profit_center,sku,brand,account,currency,actual,budget,forecast\n'


def demo_data():
    data = merge_datasets(load_csv(sample_csv()), load_csv(sample_pl_csv(), 'Sales'))
    return apply_products(data, load_product_master(sample_product_master_csv()))[0]


class CentreTests(unittest.TestCase):
    def test_profit_centre_column_is_the_centre(self):
        rows = load_csv(HEADER + '2026-08,India,PC-IN,AUR-100,Aurora,Product sales,INR,10,9,9\n', 'Sales')
        self.assertEqual((rows.dims, rows[0][CC]), (('cost_center', 'sku', 'brand'), 'PC-IN'))

    def test_a_file_with_both_keeps_the_profit_centre_as_a_level(self):
        text = ('month,country,cost_center,profit_center,account,currency,actual,budget,forecast\n'
                '2026-08,India,IN-OPS,PC-IN,Travel,INR,10,9,9\n')
        rows = load_csv(text)
        self.assertEqual((rows[0][CC], rows[0]['profit_center']), ('IN-OPS', 'PC-IN'))

    def test_upload_screen_recognises_centres_and_products(self):
        roles = guess_roles(['Month', 'Profit Center', 'Cost Center', 'Material Code', 'Brand', 'Category', 'GL Account',
                             'Actual', 'Budget', 'Forecast'], 'pl')
        self.assertEqual(roles[1:6], ['cost_center', 'extra', 'extra', 'extra', 'extra'])


class LevelTests(unittest.TestCase):
    def test_product_may_be_blank_but_other_levels_may_not(self):
        rows = load_csv(HEADER + '2026-08,India,IN-COR,,,Income tax,INR,10,9,9\n')
        self.assertEqual(rows[0]['dims'], ('cost_center',))
        with self.assertRaisesRegex(ValueError, 'missing value'):
            load_csv('month,country,project,account,currency,actual,budget,forecast\n2026-08,India,,Travel,INR,1,1,1\n')

    def test_costs_and_sales_keep_their_own_levels(self):
        rows = aggregate(demo_data(), (CC, 'brand'))
        sales = next(r for r in rows if r['account'] == 'Product sales' and r.get('brand') == 'Nimbus' and r['country'] == 'India')
        travel = next(r for r in rows if r['account'] == 'Travel' and r[CC] == 'IN-SAL')
        self.assertEqual((sales['dims'], sales['variance'] if 'variance' in sales else sales['actual'] - sales['budget']),
                         (('cost_center', 'brand'), D(-360000)))
        self.assertEqual(travel['dims'], ('cost_center',))
        self.assertEqual(short_name({'dims': {CC: 'PC-IN', 'brand': 'Nimbus'}, 'attrs': {'cc_name': 'India Commercial'},
                                     'country': 'India', 'account': 'Product sales'}), 'India Commercial Nimbus Product sales')

    def test_brand_lines_needing_explanation(self):
        lines = review_lines(aggregate(demo_data(), (CC, 'brand')), 'budget', (D(25000), D(10)), (D(50000), D(5)))
        brand = {(l['brand'], l['account']): l['variance'] for l in lines if l['material'] and l.get('brand')}
        self.assertEqual(brand, {('Nimbus', 'Product sales'): D(-360000), ('Aurora', 'Cost of goods sold'): D(300000)})


class ProductMasterTests(unittest.TestCase):
    def test_master_rules(self):
        self.assertEqual(load_product_master('SKU,Brand Name,Product Category\nA1,Aurora,Care\n'),
                         {'A1': {'brand': 'Aurora', 'category': 'Care'}})
        for text, message in (('sku,name\nA1,x\n', 'brand or category column'), ('sku,brand\nA1,X\nA1,X\n', 'listed twice'),
                              ('sku,brand\nA1,\n', 'no brand or category'),
                              ('sku,brand,category\nA1,X,C1\nA2,X,C2\n', 'two categories')):
            with self.subTest(text=text), self.assertRaisesRegex(ValueError, message):
                load_product_master(text)

    def test_sku_data_gets_brand_and_category(self):
        data = demo_data()
        self.assertEqual(data.dims, ('cost_center', 'category', 'brand', 'sku'))
        row = next(r for r in data if r.get('sku') == 'NIM-100')
        self.assertEqual((row['brand'], row['category'], row['sku_name']), ('Nimbus', 'Personal care', 'Nimbus Soap 4-pack'))
        self.assertNotIn('brand', next(r for r in data if r['account'] == 'Income tax'))

    def test_unknown_sku_stops_with_a_template(self):
        data = load_csv(HEADER + '2026-08,India,PC-IN,ZZZ-9,,Product sales,INR,10,9,9\n', 'Sales')
        with self.assertRaises(UnmappedSKU) as caught:
            apply_products(data, load_product_master(sample_product_master_csv()))
        self.assertEqual(caught.exception.codes, ['ZZZ-9'])
        text = table_csv(product_workbook({}, ['007'], ['007']), 'p.xlsx')
        self.assertIn('007,,,,New: fill in the brand and category', text)  # leading zeros survive

    def test_contradictions_stop(self):
        data = load_csv(HEADER + '2026-08,India,PC-IN,AUR-100,Nimbus,Product sales,INR,10,9,9\n', 'Sales')
        with self.assertRaisesRegex(ValueError, 'the data says brand Nimbus, the product master says Aurora'):
            apply_products(data, load_product_master(sample_product_master_csv()))
        two = load_csv(HEADER + '2026-08,India,PC-IN,A1,Aurora,Product sales,INR,1,1,1\n'
                                '2026-08,India,PC-SG,A1,Nimbus,Product sales,INR,1,1,1\n', 'Sales')
        with self.assertRaisesRegex(ValueError, 'SKU A1 is in two brands'):
            check_products(two)

    def test_brand_only_sales_need_no_master(self):
        text = HEADER.replace('sku,', '') + '2026-08,India,PC-IN,Aurora,Product sales,INR,110,100,100\n'
        rows = aggregate(check_products(load_csv(text, 'Sales')), (CC, 'brand'))
        self.assertEqual((rows[0]['brand'], rows[0]['dims']), ('Aurora', ('cost_center', 'brand')))


class ProductOutputTests(unittest.TestCase):
    def test_sku_drill_down_adds_up(self):
        data = demo_data()
        rows = aggregate(data, (CC, 'brand'))
        line = next(r for r in rows if r.get('brand') == 'Nimbus' and r['country'] == 'India' and r['account'] == 'Product sales')
        parts = breakdown(data, line, 'budget', 'sku')
        self.assertEqual([(p['sku'], p['actual'] - p['budget']) for p in parts], [('NIM-100', D(-300000)), ('NIM-200', D(-60000))])
        self.assertEqual(sum(p['actual'] for p in parts), line['actual'])

    def test_margins_by_brand(self):
        margins = {m['label']: m for m in product_margins(demo_rows(), 'budget', 'brand')}
        aurora = margins['Aurora']
        self.assertEqual((aurora['sales_variance'], aurora['gm_variance']), (D(350000), D(20000)))  # margin INR up, margin % down
        self.assertEqual((round(aurora['gm_pct_baseline'], 1), round(aurora['gm_pct_actual'], 1), round(aurora['gm_pts'], 2)),
                         (D('55.0'), D('53.6'), D('-1.40')))
        self.assertEqual(sum(m['sales_actual'] for m in margins.values()), D(30500000))  # every sales line has a brand

    def test_pack_excel_and_deck(self):
        pack = lock_pack(demo_pack())
        self.assertEqual(pack['products']['level'], 'brand')
        book = load_workbook(io.BytesIO(pack_excel(pack)))
        self.assertIn('Margins by brand', book.sheetnames)
        self.assertIn('By brand', book.sheetnames)
        self.assertEqual(products_headline(pack), 'Breeze adds INR 410,000 of sales; Nimbus loses INR 260,000. Mix moves inside the totals.')
        _, texts = slide_texts(make_deck(pack, None)['pptx'])
        self.assertIn('Aurora\n12,350,000\n+350,000\n2.9%\n6,620,000\n53.6%\n−1.4 pts', texts[4])

    def test_sku_spellings_are_offered_as_a_match(self):
        roles = ['month', 'country', 'sku', 'account', 'amount']
        specs = {'actuals': (['M', 'C', 'SKU', 'A', 'Amt'], [['2026-08', 'India', 'AUR-100', 'Sales', 1]], roles, 'a.csv'),
                 'budget': (['M', 'C', 'SKU', 'A', 'Amt'], [['2026-08', 'India', 'aur100', 'Sales', 1]], roles, 'b.csv')}
        _, _, suggestions = validate.check(specs)
        self.assertEqual([(s['field'], s['to']) for s in suggestions], [('sku', 'AUR-100')])


class AppTests(unittest.TestCase):
    def test_brand_view_drill_down_and_sku_level(self):
        app = AppTest.from_file(str(APP), default_timeout=60).run()
        self.assertFalse(app.exception)
        level = next(s for s in app.selectbox if s.label == 'Sales and cost of sales by')
        self.assertEqual((level.value, level.options), ('brand', ['SKU', 'Brand', 'Category', 'No product split']))
        self.assertTrue(any('Brand' in d.value.columns and 'Gross margin %' in d.value.columns for d in app.dataframe))
        self.assertTrue(any(e.label == 'SKUs in this line (2)' for e in app.expander))  # the Nimbus line comes first
        level.set_value('sku').run()
        self.assertFalse(app.exception)
        self.assertEqual(next(m.value for m in app.metric if m.label == 'Net profit · INR'), '2,270,000')
        self.assertFalse(any(e.label.startswith('SKUs in this line') for e in app.expander))

    def test_unknown_sku_upload_offers_the_master(self):
        app = AppTest.from_file(str(APP), default_timeout=60).run()
        app.sidebar.file_uploader[0].upload('expenses.csv', sample_csv().encode(), 'text/csv').run()
        sales = sample_pl_csv().replace('BRZ-200', 'BRZ-300')
        app.sidebar.file_uploader[1].upload('sales.csv', sales.encode(), 'text/csv').run()
        app.sidebar.file_uploader[4].upload('products.csv', sample_product_master_csv().encode(), 'text/csv').run()
        self.assertTrue(any('not in the product master: BRZ-300' in e.value for e in app.error))


if __name__ == '__main__':
    unittest.main()
