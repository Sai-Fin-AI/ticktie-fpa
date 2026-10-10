"""CFO deck agent: slides, charts and tables from the pack; AI headlines tick-and-tied or replaced."""
import io
import sys
import unittest
from decimal import Decimal as D
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu
from streamlit.testing.v1 import AppTest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))

from agents.deck import KEYS, account_steps, default_headlines, make_deck, owner_table, timing_split  # noqa: E402
from checker import misattributed_figures  # noqa: E402
from finance import load_csv, sample_csv  # noqa: E402
from llm import DEMO, AIError, Client  # noqa: E402
from pack import build_pack, lock_pack, pack_json  # noqa: E402
from record_demo import demo_pack  # noqa: E402
from test_agents import FakeClient  # noqa: E402

APP = ROOT / 'app.py'


def slide_texts(pptx):
    deck = Presentation(io.BytesIO(pptx))
    out = []
    for slide in deck.slides:
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            if shape.has_table:
                parts += [c.text for row in shape.table.rows for c in row.cells]
        out.append('\n'.join(parts))
    return deck, out


def good_headlines():
    return {'headlines': {'summary': 'Expenses INR 365,000 above budget; 66% explained',
                          'bridge': 'Distribution adds INR 340,000 to the budget-to-actual bridge',
                          'top': 'India Operations Distribution leads at INR 240,000',
                          'heatmap': 'Overspend concentrates in Distribution and Consulting',
                          'offsets': 'Singapore Employee costs nets to INR 0 but moves INR 180,000 gross',
                          'owners': 'Priya Nair holds INR 170,000 of open variance',
                          'outlook': 'Recurring INR 390,000 needs a forecast update'},
            'summary': ['Actual INR 12,375,000 vs budget INR 12,010,000.', 'INR 575,000 remains unexplained.',
                        'Distribution is the largest account movement.']}


class DeckFigureTests(unittest.TestCase):
    def setUp(self):
        self.pack = lock_pack(demo_pack())

    def test_derived_figures_reconcile(self):
        steps = dict(account_steps(self.pack))  # P&L demo: the default bridge is EBITDA, in profit terms
        self.assertEqual(steps, {'Product sales': D('500000'), 'Cost of goods sold': D('-400000'), 'Distribution': D('-340000'),
                                 'Consulting': D('-130000'), 'Employee costs': D('50000'), 'Warehousing': D('40000'),
                                 'Recruitment': D('15000')})
        self.assertEqual(sum(steps.values()), D('-265000'))
        split = timing_split(self.pack)
        self.assertEqual((split['One-off'], split['Recurring'], split['Timing difference'], split['Unexplained']),
                         (D('570000'), D('390000'), D('150000'), D('575000')))
        self.assertEqual(sum(v for k, v in split.items() if k != 'Unexplained'), D('1110000'))
        self.assertEqual(owner_table(self.pack)[0][0], 'Priya Nair')
        self.assertEqual(owner_table(self.pack)[0][1]['open'], D('170000'))

    def test_binding_check_catches_figures_on_the_wrong_line(self):
        from agents.deck import derived_records
        extra = derived_records(self.pack)
        self.assertTrue(misattributed_figures('India Distribution is INR 340,000 above budget.', self.pack, extra))
        for ok in ('Distribution adds INR 340,000', 'India Operations Distribution leads at INR 240,000',
                   'Priya Nair holds INR 170,000', 'Recurring INR 390,000', 'Total INR 365,000 above budget'):
            with self.subTest(ok=ok):
                self.assertEqual(misattributed_figures(ok, self.pack, extra), [])


class DeckBuildTests(unittest.TestCase):
    def setUp(self):
        self.pack = lock_pack(demo_pack())

    def test_slides_charts_and_tables(self):
        result = make_deck(self.pack, FakeClient(good_headlines()))
        self.assertEqual((result['issues'], result['source']), ([], 'AI (Fake model), tick-and-tied'))
        deck, texts = slide_texts(result['pptx'])
        self.assertEqual(len(deck.slides), 12)
        self.assertAlmostEqual(Emu(deck.slide_width).inches, 13.333, places=2)
        self.assertEqual(result['slides'], 12)
        self.assertIn('P&L variance review\nAugust 2026', texts[0])
        self.assertIn(self.pack['content_hash'][:12], texts[0])
        self.assertIn('Net profit INR 195,000 below budget; EBITDA INR 265,000 below', texts[2])
        self.assertIn('GROSS MARGIN', texts[2])
        self.assertIn('Distribution adds INR 340,000', texts[3])
        self.assertIn('EBITDA BRIDGE: BUDGET TO ACTUAL', texts[3])
        charts = [s for s in deck.slides[3].shapes if s.has_chart]
        self.assertEqual(list(charts[0].chart.plots[0].categories),
                         ['Budget', 'Product sales', 'Cost of goods sold', 'Distribution', 'Consulting', 'Employee costs',
                          'Warehousing', 'Recruitment', 'Actual'])
        self.assertIn('SALES AND GROSS MARGIN BY BRAND', texts[4])
        self.assertIn('Breeze adds INR 410,000 of sales; Nimbus loses INR 260,000', texts[4])
        self.assertIn('Arjun Mehta', texts[5])
        self.assertIn('+240k', texts[6])
        self.assertIn('India | Product sales\nNet +300,000, gross 1,020,000: offsetting lines hide 720,000.', texts[7])  # brand mix
        self.assertIn('Priya Nair', texts[8])
        self.assertIn('Recurring INR 390,000', texts[9])
        self.assertIn('Lines needing explanation (1–12 of 16)', texts[10])
        self.assertIn('(13–16 of 16)', texts[11])
        self.assertEqual([k for k, _ in result['outline']][:4], ['Title', 'Executive summary', 'Profit and loss', 'Bridge'])

    def test_bridge_follows_the_choice_on_screen(self):
        for measure, first, cats in (('Sales', 'Product sales', 3), ('Costs', 'Cost of goods sold', 9),
                                     ('Gross margin', 'Product sales', 4), ('Net profit', 'Product sales', 10)):
            with self.subTest(measure=measure):
                result = make_deck(self.pack, None, bridge=measure)
                deck, texts = slide_texts(result['pptx'])
                chart = next(s for s in deck.slides[3].shapes if s.has_chart).chart
                self.assertEqual(list(chart.plots[0].categories)[1], first)
                self.assertEqual(len(list(chart.plots[0].categories)), cats)
                self.assertEqual(result['issues'], [])
        _, texts = slide_texts(make_deck(self.pack, None, bridge='Costs')['pptx'])
        self.assertIn('Cost of goods sold adds INR 400,000 to costs', texts[3])

    def test_failing_headlines_are_replaced_not_shown(self):
        bad = good_headlines()
        bad['headlines']['top'] = 'India Distribution is INR 340,000 above budget'
        bad['headlines']['owners'] = 'Priya Nair holds INR 999,000 (see L03)'
        client = FakeClient(bad, bad)
        result = make_deck(self.pack, client)
        defaults = default_headlines(self.pack)
        outline = dict(result['outline'])
        self.assertEqual(outline['Largest movements'], defaults['top'])
        self.assertEqual(outline['Open items by owner'], defaults['owners'])
        self.assertEqual(outline['Bridge'], 'Distribution adds INR 340,000 to the budget-to-actual bridge')
        self.assertEqual(len(result['issues']), 2)
        self.assertIn('does not belong to India / Distribution', client.calls[1][1])
        _, texts = slide_texts(result['pptx'])
        self.assertNotIn('999,000', '\n'.join(texts))

    def test_ai_down_still_builds_with_standard_titles(self):
        result = make_deck(self.pack, FakeClient(AIError('Could not reach Ollama.')))
        self.assertEqual(result['source'], 'standard titles')
        self.assertIn('Could not reach Ollama', result['issues'][0])
        self.assertEqual(dict(result['outline'])['Bridge'], default_headlines(self.pack)['bridge'])

    def test_account_level_pack_skips_the_offsets_slide(self):
        pack = lock_pack(build_pack(load_csv(sample_csv('account')), 'budget', D('50000'), D('5'), {}))
        result = make_deck(pack, None)
        deck, texts = slide_texts(result['pptx'])
        self.assertNotIn('What net figures hide', [k for k, _ in result['outline']])
        self.assertEqual(len(deck.slides), 8)
        self.assertIn('India', texts[4])
        self.assertIn('+200k', texts[4])

    def test_edited_pack_is_refused_and_demo_recording_matches(self):
        import json
        tampered = json.loads(pack_json(self.pack))
        tampered['totals']['actual'] = '1'
        with self.assertRaises(ValueError):
            make_deck(tampered, None)
        result = make_deck(self.pack, Client(DEMO))
        self.assertTrue(result['source'].startswith('AI (Demo'), result['issues'])
        # The recording is found; any headline that fails the checks is swapped for standard wording, never shown.
        self.assertTrue(all('replaced with standard wording' in i for i in result['issues']), result['issues'])
        self.assertEqual(set(KEYS), set(default_headlines(self.pack)))


class DeckAppTests(unittest.TestCase):
    def test_one_click_runs_both_agents(self):
        app = AppTest.from_file(str(APP), default_timeout=60).run()
        button = lambda label: next(b for b in app.button if b.label == label)
        button('Load demo explanations').click().run()
        button('Lock review pack').click().run()
        headlines = next(r for r in app.radio if r.label == 'Deck headlines')
        self.assertEqual(headlines.value, 'Standard, written by the engine')
        headlines.set_value('AI-written, tick-and-tied').run()
        button('Draft CFO commentary and deck').click().run()
        self.assertFalse(app.exception)
        deck = app.session_state['ai_deck']
        self.assertTrue(deck['source'].startswith('AI (Demo'), deck['issues'])
        self.assertEqual(len(Presentation(io.BytesIO(deck['pptx'])).slides), 12)
        self.assertTrue(any('CFO deck' in x.value and '12 slides' in x.value for x in app.markdown))
        self.assertIn('✔ Tick-and-tie passed', app.session_state['ai_commentary']['html'])
        next(r for r in app.radio if r.label == 'Deck headlines').set_value('Standard, written by the engine').run()
        button('Draft CFO commentary and deck').click().run()
        self.assertEqual(app.session_state['ai_deck']['source'], 'standard titles')
        self.assertNotIn('error', app.session_state['ai_commentary'])

    def test_deck_only_button_needs_no_ai(self):
        app = AppTest.from_file(str(APP), default_timeout=60).run()
        button = lambda label: next(b for b in app.button if b.label.startswith(label))
        button('Lock review pack').click().run()
        button('🖥️ CFO deck only').click().run()
        self.assertFalse(app.exception)
        self.assertEqual(app.session_state['ai_deck']['source'], 'standard titles')
        self.assertNotIn('ai_commentary', app.session_state)


if __name__ == '__main__':
    unittest.main()
