"""Negative balances: accepted, flagged, always explained, never rolled forward, and shown with their sign."""
import io
import sys
import unittest
from decimal import Decimal as D
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))

import fmt  # noqa: E402
from agents.deck import default_summary, make_deck  # noqa: E402
from agents.writer import standard_line, standard_summary  # noqa: E402
from checker import allowed_figures, unsupported_figures  # noqa: E402
from evidence import make_note  # noqa: E402
from finance import aggregate, commentary, load_csv, review_lines  # noqa: E402
from pack import build_pack, lock_pack, pack_excel  # noqa: E402
from record_demo import demo_pack  # noqa: E402
from test_deck import slide_texts  # noqa: E402

HEADER = 'month,country,account,currency,actual,budget,forecast\n'
LINE = (D(50000), D(5))


def lines(text):
    return {r['account']: r for r in review_lines(aggregate(load_csv(HEADER + text), ()), 'budget', LINE)}


class RuleTests(unittest.TestCase):
    def test_small_negative_line_is_flagged_and_needs_an_explanation(self):
        found = lines('2026-08,India,Travel,INR,-15000,10000,10000\n2026-08,India,Rent,INR,12000,10000,10000\n')
        travel, rent = found['Travel'], found['Rent']
        self.assertEqual((travel['negative'], travel['material'], travel['scope']), (True, True, 'Negative balance'))
        self.assertEqual((rent['negative'], rent['material']), (False, False))
        self.assertIn('Negative balance: the actual is INR -15,000.00, usually a reversal', commentary(travel, 'budget'))

    def test_large_negative_line_keeps_its_scope_and_the_flag(self):
        travel = lines('2026-08,India,Travel,INR,-150000,100000,100000\n')['Travel']
        self.assertEqual((travel['negative'], travel['scope'], travel['direction']), (True, 'Line', 'Favorable'))
        self.assertEqual(travel['pct'], D(-250))  # divided by the size of the budget

    def test_negative_budget_is_flagged_too_and_no_movement_is_not(self):
        found = lines('2026-08,India,Recharges,INR,-30000,-50000,-50000\n2026-08,India,Recovery,INR,-5000,-5000,-5000\n')
        self.assertEqual((found['Recharges']['negative'], found['Recharges']['material']), (True, True))
        self.assertEqual((found['Recovery']['negative'], found['Recovery']['material']), (True, False))  # on plan

    def test_recurring_is_refused(self):
        travel = lines('2026-08,India,Travel,INR,-15000,10000,10000\n')['Travel']
        with self.assertRaisesRegex(ValueError, 'cannot be marked Recurring'):
            make_note(travel, 'budget', 'Credit note.', 'CN-1', '25000', 'Confirmed by reviewer', 'Recurring')
        self.assertEqual(make_note(travel, 'budget', 'Credit note.', 'CN-1', '25000', 'Confirmed by reviewer', 'One-off')['timing'],
                         'One-off')

    def test_balance_format_keeps_the_sign(self):
        self.assertEqual((fmt.balance(-15000), fmt.balance(15000), fmt.group(-15000)), ('−15,000', '15,000', '15,000'))


class OutputTests(unittest.TestCase):
    def setUp(self):
        self.pack = lock_pack(demo_pack())
        self.line = next(l for l in self.pack['lines'] if l['account'] == 'Recruitment')

    def test_pack_marks_the_line(self):
        self.assertEqual((self.line['negative'], self.line['scope'], self.line['actual']), (True, 'Negative balance', '-5000'))
        self.assertNotIn('negative', next(l for l in self.pack['lines'] if l['account'] == 'Travel'))
        book = load_workbook(io.BytesIO(pack_excel(self.pack)))
        row = next(r for r in book['Variances'].iter_rows(min_row=2, values_only=True) if r[2] == 'Recruitment')
        self.assertEqual((row[4], row[-1]), (-5000, 'Yes (negative balance)'))

    def test_standard_commentary_states_it_and_ties_out(self):
        text = standard_line(self.pack, self.line)
        self.assertIn('negative balance of INR −5,000.00', text)
        self.assertEqual(unsupported_figures(text, allowed_figures(self.pack)), [])

    def test_deck_shows_the_sign_and_the_flag(self):
        _, texts = slide_texts(make_deck(self.pack, None)['pptx'])
        appendix = '\n'.join(texts[9:])
        self.assertIn('India | IN-COR India Corporate | Recruitment · ⚠ negative balance\n−5,000\n10,000\n−15,000', appendix)


class CheckerTests(unittest.TestCase):
    def test_minus_signed_figures_are_checked(self):
        allowed = allowed_figures(lock_pack(demo_pack()))
        self.assertEqual(unsupported_figures('The actual is INR -5,000.00.', allowed), [])
        self.assertTrue(unsupported_figures('The actual is INR -7,777.00.', allowed))  # was skipped before
        self.assertEqual(unsupported_figures('Ref CN-0731, FY2026-27, Aug-2026.', allowed), [])  # codes still ignored


class LossTests(unittest.TestCase):
    def test_a_loss_keeps_its_minus_sign(self):
        text = (HEADER.replace('account,', 'account,pl_line,') + '2026-08,India,Product sales,Sales,INR,100000,200000,200000\n'
                '2026-08,India,Rent,Operating expenses,INR,150000,120000,120000\n')
        pack = lock_pack(build_pack(aggregate(load_csv(text), ()), 'budget', D(1000), D(1), {}))
        self.assertEqual(pack['totals']['actual'], '-50000')  # net profit: a loss
        self.assertIn('Net profit INR −50,000 vs budget INR 80,000', default_summary(pack)[0])
        self.assertIn('actual INR −50,000.00 against INR 80,000.00', standard_summary(pack)[0])


if __name__ == '__main__':
    unittest.main()
