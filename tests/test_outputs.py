"""Number styles, full-year view, line naming, reason column and bulk explanation entry."""
import io
import json
import sys
import unittest
from decimal import Decimal as D
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))

import fmt  # noqa: E402
from agents.deck import make_deck  # noqa: E402
from checker import account_total_confusion, misattributed_figures  # noqa: E402
from evidence import note_key  # noqa: E402
from finance import load_csv, review_lines, sample_csv  # noqa: E402
from notes_panel import apply_rows, explanation_template, read_template  # noqa: E402
from pack import build_pack, fy_outlook, lock_pack, pack_excel  # noqa: E402
from record_demo import LINE_THRESHOLDS, TOTAL_THRESHOLDS, demo_notes, demo_pack, demo_rows  # noqa: E402
from test_deck import slide_texts  # noqa: E402


class NumberStyleTests(unittest.TestCase):
    def tearDown(self):
        fmt.use(None)

    def test_international_and_indian(self):
        fmt.use('International (k / m)')
        self.assertEqual((fmt.group(12380000), fmt.compact(340000), fmt.compact(-12380000)), ('12,380,000', '+340k', '−12.38m'))
        fmt.use('Indian (lakh / crore)')
        self.assertEqual(fmt.group(12380000), '1,23,80,000')
        self.assertEqual(fmt.group(380000, 2), '3,80,000.00')
        self.assertEqual((fmt.group(999), fmt.group(1000), fmt.group(100000)), ('999', '1,000', '1,00,000'))
        self.assertEqual((fmt.compact(340000), fmt.compact(-12380000), fmt.compact(40000)), ('+3.40 L', '−1.24 Cr', '+0.40 L'))

    def test_indian_deck_and_checker_agree(self):
        pack = lock_pack(demo_pack())
        result = make_deck(pack, None, 'Indian (lakh / crore)')
        self.assertEqual(result['issues'], [])
        _, texts = slide_texts(result['pptx'])
        self.assertIn('+3,80,000', texts[0])
        self.assertIn('+2.40 L', texts[4])
        self.assertIn('INR crores', texts[2])
        # International style is back for the next caller in this thread only if asked for.
        fmt.use(None)
        self.assertEqual(fmt.group(380000), '380,000')


class AppSettingsTests(unittest.TestCase):
    def test_indian_style_and_financial_year_on_screen(self):
        from streamlit.testing.v1 import AppTest
        app = AppTest.from_file(str(ROOT / 'app.py'), default_timeout=60).run()
        next(b for b in app.button if b.label == 'Load demo explanations').click().run()
        caption = lambda: next(c.value for c in app.caption if 'Full-year view' in c.value)
        self.assertIn('FY2026-27, 7 months left', caption())
        next(s for s in app.selectbox if s.label == 'Number style').set_value('Indian (lakh / crore)').run()
        self.assertEqual(next(m for m in app.metric if m.label == 'Actual expense · INR').value, '1,23,80,000')
        self.assertIn('INR −4,65,000', caption())
        next(s for s in app.selectbox if s.label == 'Financial year starts in').set_value('January').run()
        self.assertFalse(app.exception)
        self.assertIn('FY2026, 4 months left', caption())
        self.assertEqual([t.label for t in app.tabs][:3], ['📝 One line at a time', '📋 All lines (13, all countries)',
                                                         '📤 Excel round-trip'])


class FullYearTests(unittest.TestCase):
    def test_outlook_for_the_demo(self):
        outlook = lock_pack(demo_pack())['outlook']
        self.assertEqual((outlook['fy'], outlook['months_left'], outlook['fy_start']), ('FY2026-27', 7, 'April'))
        self.assertEqual({k: D(outlook[k]) for k in ('one_off', 'recurring_month', 'recurring_fy', 'timing', 'fy_effect')},
                         {'one_off': D(255000), 'recurring_month': D(-90000), 'recurring_fy': D(-720000),
                          'timing': D(-150000), 'fy_effect': D(-465000)})

    def test_calendar_year_and_last_month(self):
        rows = demo_rows()
        lines = build_pack(rows, 'budget', *LINE_THRESHOLDS, demo_notes(rows), TOTAL_THRESHOLDS)['lines']
        self.assertEqual(fy_outlook(lines, '2026-08', 1)['months_left'], 4)
        self.assertEqual(fy_outlook(lines, '2026-08', 1)['fy'], 'FY2026')
        self.assertEqual(fy_outlook(lines, '2027-03', 4)['months_left'], 0)
        self.assertEqual(fy_outlook(lines, '2027-03', 4)['fy'], 'FY2026-27')

    def test_full_year_in_deck_excel_and_checker(self):
        pack = lock_pack(demo_pack())
        _, texts = slide_texts(make_deck(pack, None)['pptx'])
        outlook_slide = texts[7]
        self.assertIn('FULL-YEAR VIEW · FY2026-27 · 7 MONTHS LEFT', outlook_slide)
        self.assertIn('−465,000', outlook_slide)
        self.assertIn('Confirmed items point to INR 465,000 below plan for FY2026-27', outlook_slide)
        summary = load_workbook(io.BytesIO(pack_excel(pack)))['Summary']
        values = {row[0]: row[1] for row in summary.iter_rows(values_only=True)}
        self.assertEqual(values['Full-year view: effect of confirmed explanations'], -465000)
        self.assertEqual(misattributed_figures('Recurring items add INR 720,000 for the full year.', pack), [])


class LineNamingTests(unittest.TestCase):
    def setUp(self):
        self.pack = lock_pack(demo_pack())
        self.line = next(l for l in self.pack['lines'] if l['dims'].get('cost_center') == 'IN-OPS' and l['account'] == 'Distribution')

    def test_account_total_confusion(self):
        self.assertTrue(account_total_confusion('India Distribution was INR 240,000.00 above budget.', self.line, self.pack))
        self.assertTrue(account_total_confusion('India | Distribution rose.', self.line, self.pack))
        for ok in ('India Operations Distribution was INR 240,000.00 above budget.', 'IN-OPS Distribution rose.',
                   'Urgent air freight explains INR 180,000.00.'):
            with self.subTest(ok=ok):
                self.assertEqual(account_total_confusion(ok, self.line, self.pack), [])

    def test_account_total_figure_must_be_the_total(self):
        self.assertTrue(misattributed_figures('India | Distribution was INR 240,000.00 above budget.', self.pack))
        self.assertEqual(misattributed_figures('India | Distribution was INR 200,000.00 above budget.', self.pack), [])
        self.assertEqual(misattributed_figures('India Operations Distribution was INR 240,000.00 above budget.', self.pack), [])

    def test_each_amount_binds_to_the_name_before_it(self):
        right = ('India Operations Distribution (INR 180,000.00 explained) and '
                 'Australia Sales Employee costs (INR 150,000.00 favorable).')
        swapped = ('India Operations Distribution (INR 70,000.00 explained) and '  # 70,000 is another line's figure
                   'Australia Sales Employee costs (INR 180,000.00 favorable).')
        self.assertEqual(misattributed_figures(right, self.pack), [])
        self.assertEqual(len(misattributed_figures(swapped, self.pack)), 2)
        self.assertEqual(misattributed_figures('Of the INR 380,000 net, India Operations Distribution moved most.', self.pack), [])

    def test_account_level_data_is_unaffected(self):
        pack = lock_pack(build_pack(load_csv(sample_csv('account')), 'budget', D('50000'), D('5'), {}))
        line = next(l for l in pack['lines'] if l['material'])
        self.assertEqual(account_total_confusion(f"{line['country']} {line['account']} moved.", line, pack), [])

    def test_reason_column_on_largest_movements(self):
        _, texts = slide_texts(make_deck(self.pack, None)['pptx'])
        self.assertIn('REASON GIVEN', texts[3])
        self.assertIn('Urgent air freight for festive-season launch stock', texts[3])
        self.assertIn('India Operations Distribution\nIndia · IN-OPS · Arjun Mehta', texts[3])
        self.assertIn('Pending from owner', texts[3])


class BulkNotesTests(unittest.TestCase):
    def setUp(self):
        self.rows = demo_rows()
        material = [r for r in review_lines(self.rows, 'budget', LINE_THRESHOLDS, TOTAL_THRESHOLDS) if r['material']]
        self.options = {note_key(r, 'budget'): r for r in material}
        self.notes = demo_notes(self.rows)

    def records(self, notes):
        return read_template(explanation_template(self.options, notes, '2026-08', 'budget'))

    def test_template_round_trip_changes_nothing(self):
        records = self.records(self.notes)
        self.assertEqual(len(records), len(self.options))
        updated, errors, changed = apply_rows(records, self.options, 'budget', self.notes)
        self.assertEqual((errors, changed), ([], 0))
        self.assertEqual(updated, self.notes)

    def test_filled_template_saves_good_rows_and_reports_bad_ones(self):
        book = load_workbook(io.BytesIO(explanation_template(self.options, {}, '2026-08', 'budget')))
        ws = book['Explanations']
        header = [c.value for c in ws[1]]
        col = {name: header.index(name) + 1 for name in header}
        lines = {ws.cell(r, col['Line']).value: r for r in range(2, ws.max_row + 1)}
        good = lines['India | IN-SAL India Sales | Travel']
        bad = lines['Singapore | SG-OPS Singapore Operations | Warehousing']
        for row, amount in ((good, 70000), (bad, 999999)):
            ws.cell(row, col['Explanation']).value = 'Fewer client visits in August.'
            ws.cell(row, col['Reference']).value = 'Fictional travel report'
            ws.cell(row, col['Amount explained (INR)']).value = amount
            ws.cell(row, col['Review status']).value = 'confirmed by reviewer'  # case does not matter
            ws.cell(row, col['Timing']).value = 'One-off'
        out = io.BytesIO()
        book.save(out)
        updated, errors, changed = apply_rows(read_template(out.getvalue()), self.options, 'budget', {})
        self.assertEqual(changed, 1)
        self.assertEqual(len(errors), 1)
        self.assertIn('Warehousing', errors[0])
        self.assertIn('cannot exceed', errors[0])
        note = next(iter(updated.values()))
        self.assertEqual((note['amount'], note['status']), (D('70000'), 'Confirmed by reviewer'))

    def test_blank_row_removes_a_note_and_edited_key_is_refused(self):
        records = self.records(self.notes)
        target = next(r for r in records if r['Explanation'])
        target['Explanation'] = target['Reference'] = ''
        updated, errors, changed = apply_rows(records, self.options, 'budget', self.notes)
        self.assertEqual((errors, changed, len(updated)), ([], 1, len(self.notes) - 1))
        book = load_workbook(io.BytesIO(explanation_template(self.options, {}, '2026-08', 'budget')))
        book['Explanations'].cell(2, 1).value = 'edited'
        out = io.BytesIO()
        book.save(out)
        with self.assertRaises(ValueError):
            read_template(out.getvalue())

    def test_other_workbooks_are_refused(self):
        from ingest import sample_budget_xlsx
        with self.assertRaises(ValueError):
            read_template(sample_budget_xlsx())
        self.assertTrue(json.loads(self.records(self.notes)[0]['Line key (do not edit)']))


if __name__ == '__main__':
    unittest.main()
