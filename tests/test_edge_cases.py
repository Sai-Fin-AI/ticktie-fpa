"""Edge cases found by the synthetic stress run (random datasets pushed through the whole pipeline)."""
import io
import os
import sys
import unittest
from decimal import Decimal as D
from pathlib import Path

from pptx import Presentation
from streamlit.testing.v1 import AppTest

import validate
from agents.deck import _pts, make_deck
from checker import wrong_claims
from agents.writer import to_markdown
from finance import aggregate, load_csv, merge_datasets, sample_pl_csv, statement
from ingest import guess_roles, normalize_month, separate_csv
from pack import build_pack, lock_pack

APP = Path(__file__).resolve().parents[1] / 'app.py'
sys.path.insert(0, str(APP.parent / 'scripts'))  # record_demo
BRANDS = 'month,country,profit_center,brand,account,currency,actual,budget,forecast\n'


def pack_of(text, default_line='Sales'):
    data = load_csv(text, default_line)
    return lock_pack(build_pack(aggregate(data, data.dims), 'budget', D(25000), D(10), {}, (D(50000), D(5))))


class SalesWithoutCostsTests(unittest.TestCase):
    """A P&L file with sales and cost of sales only: no operating expenses, so nothing sits just above EBITDA."""

    def test_statement_still_has_ebitda(self):
        rows = aggregate(load_csv(sample_pl_csv(), 'Sales'), ())
        labels = [e['label'] for e in statement(rows, 'budget') if e['kind'] == 'subtotal']
        self.assertEqual(labels[:2], ['Gross margin', 'EBITDA'])

    def test_deck_and_app(self):
        make_deck(pack_of(sample_pl_csv()), None)  # used to fail on the missing EBITDA tile
        app = AppTest.from_file(str(APP), default_timeout=60).run()
        app.sidebar.file_uploader[0].upload('sales.csv', sample_pl_csv().encode(), 'text/csv').run()
        self.assertFalse(app.exception)
        self.assertEqual({m.label: m.value for m in app.metric}['EBITDA · INR'], '16,600,000')


class HostedTests(unittest.TestCase):
    def test_online_demo_hides_the_local_model_and_warns_about_data(self):
        os.environ['FPA_HOSTED'] = '1'
        try:
            app = AppTest.from_file(str(APP), default_timeout=60).run()
        finally:
            del os.environ['FPA_HOSTED']
        self.assertFalse(app.exception)
        provider = next(s for s in app.sidebar.selectbox if s.label == 'AI provider')
        self.assertNotIn('Ollama (this computer)', provider.options)
        self.assertTrue(any('never company data' in c.value for c in app.caption))
        local = AppTest.from_file(str(APP), default_timeout=60).run()
        self.assertIn('Ollama (this computer)', next(s for s in local.sidebar.selectbox if s.label == 'AI provider').options)


class InputShapeTests(unittest.TestCase):
    def test_names_with_a_pipe_are_not_a_false_duplicate(self):
        main = load_csv('month,country,cost_center,account,currency,actual,budget,forecast\n'
                        '2026-08,India,"A | B",Travel,INR,1,1,1\n')
        extra = load_csv(BRANDS + '2026-08,India,A,B,Travel,INR,1,1,1\n')
        self.assertEqual(len(merge_datasets(main, extra)), 2)

    def test_separate_files_may_leave_the_material_blank(self):
        headers = ['Period', 'Entity', 'Profit Center', 'Material', 'GL Account', 'Amount']
        roles = guess_roles(headers, 'actuals')
        actuals = (headers, [['Aug-2026', 'India', 'PC-IN', 'M1', 'Product sales', 500],
                             ['Aug-2026', 'India', 'PC-IN', '', 'Travel', 20]], roles, 'a.csv')
        budget = (headers, [['2026-08', 'India', 'PC-IN', 'M1', 'Product sales', 450],
                            ['2026-08', 'India', 'PC-IN', '', 'Travel', 25]], roles, 'b.csv')
        text, _, report = separate_csv({'actuals': actuals, 'budget': budget})
        rows = load_csv(text, 'Sales')
        self.assertEqual([(r.get('sku'), r['actual'], r['budget']) for r in rows],
                         [('M1', D(500), D(450)), (None, D(20), D(25))])
        self.assertEqual(report['unbudgeted'], [])

    def test_sap_period_without_separator(self):
        self.assertEqual((normalize_month('202608'), normalize_month(202608.0)), ('2026-08', '2026-08'))

    def test_excel_number_saved_as_text_is_offered_as_a_match(self):
        roles = ['month', 'country', 'account', 'amount']
        specs = {'actuals': (['M', 'C', 'A', 'Amt'], [['2026-08', 'India', 510100, 1]], roles, 'a.csv'),
                 'budget': (['M', 'C', 'A', 'Amt'], [['2026-08', 'India', '510100.0', 1], ['2026-08', 'India', '0510 100', 1]],
                            roles, 'b.csv')}
        _, _, suggestions = validate.check(specs)
        self.assertEqual([(v['value'], v['difference']) for v in suggestions[0]['from']],
                         [('0510 100', 'separators (- . _ / or spaces) and leading zeros'),
                          ('510100.0', 'a trailing .0 (Excel number format)')])


class SecondRoundTests(unittest.TestCase):
    """Found in the second critical round: formula injection, a deck that grew with the ledger, SKU spellings."""

    def test_text_never_becomes_a_formula_in_exports(self):
        from openpyxl import load_workbook
        from agents.interviewer import interview, query_list_csv
        from gl_panel import mapping_workbook, product_workbook
        from pack import pack_excel
        evil = '"=HYPERLINK(""http://evil.example"",""Click"")"'
        pack = pack_of('month,country,cost_center,account,currency,actual,budget,forecast\n'
                       f'2026-08,India,=cmd,{evil},INR,200000,100000,100000\n', 'Operating expenses')
        for data in (pack_excel(pack), mapping_workbook({}, ['=1+2'], ['=1+2']), product_workbook({}, ['@x'], ['@x'])):
            book = load_workbook(io.BytesIO(data))
            self.assertEqual([c.coordinate for ws in book for row in ws.iter_rows() for c in row if c.data_type == 'f'], [])
        self.assertIn('=HYPERLINK', load_workbook(io.BytesIO(pack_excel(pack)))['Variances']['C2'].value)  # kept as text
        findings, _ = interview(pack, None)
        rows = query_list_csv(findings).splitlines()[1:]
        self.assertTrue(rows and all(not cell.startswith(('=', '+', '-', '@')) for row in rows for cell in row.split(',')))

    def test_deck_appendix_stops_at_sixty_lines(self):
        text = 'month,country,cost_center,account,currency,actual,budget,forecast\n' + ''.join(
            f'2026-08,India,C{i},Travel,INR,{200000 + i},100000,100000\n' for i in range(90))
        slides = Presentation(io.BytesIO(make_deck(pack_of(text, 'Operating expenses'), None)['pptx'])).slides
        titles = [' '.join(s.text_frame.text for s in slide.shapes if s.has_text_frame) for slide in slides]
        appendix = [t for t in titles if 'Lines needing explanation' in t]
        self.assertEqual(len(appendix), 5)  # 60 of 90 lines; the Excel pack lists all of them
        self.assertIn('all 90 are in the Excel pack', appendix[-1])

    def test_sku_columns_and_the_product_master_are_spell_checked(self):
        headers = ['Month', 'Country', 'Profit Center', 'SKU', 'Account', 'Actual', 'Budget', 'Forecast']
        roles = guess_roles(headers, 'pl')
        self.assertEqual(roles[3], 'extra')  # the app maps SKU as an extra level, not a 'sku' role
        specs = {'pl': (headers, [['2026-08', 'India', 'PC-IN', '00123', 'Sales', 1, 1, 1]], roles, 's.csv')}
        _, _, suggestions = validate.check(specs, skus=['123'])
        self.assertEqual([(s['field'], s['to'], s['to_files']) for s in suggestions], [('sku', '123', ['Product master'])])
        fixed = validate.apply(specs, suggestions, [suggestions[0]['id']])
        self.assertEqual(fixed['pl'][1][0][3], '123')


class ClaimTests(unittest.TestCase):
    """Real figures placed in the wrong list: found in the recorded demo commentary, now rejected."""

    def setUp(self):
        from record_demo import demo_pack
        self.pack = lock_pack(demo_pack())

    def test_a_line_must_fit_its_list(self):
        favorable = ('Favorable variances include: Singapore Operations Employee costs (INR 90,000.00 below budget) and '
                     'India Corporate Tax Advisory (INR 75,000.00 above budget).')
        unexplained = 'Unexplained variances include: India Sales Employee costs (INR 60,000.00) and India Corporate Consulting (INR 75,000.00).'
        self.assertEqual(wrong_claims(favorable, self.pack), ['lists India Corporate as favorable, but it is unfavorable'])
        self.assertEqual(wrong_claims(unexplained, self.pack), ['lists India Corporate / Consulting as unexplained, but it is fully explained'])
        self.assertEqual(wrong_claims('Unfavorable variances include: India Corporate Consulting (INR 75,000.00 above budget).', self.pack), [])
        self.assertEqual(wrong_claims('India Operations Distribution (INR 240,000.00 below budget) was freight.', self.pack),
                         ['says India Operations / Distribution is below budget, but it is the other way'])

    def test_recorded_demo_summary_is_true_and_opens_with_the_result(self):
        from agents.writer import write_commentary
        from llm import Client
        bullets = write_commentary(self.pack, Client('Demo (recorded responses)'))['reply']['executive_summary']
        self.assertTrue(bullets[0].startswith('Net profit was INR 195,000.00 below budget'))
        self.assertFalse(any('Tax Advisory' in b or b.startswith('Unexplained') for b in bullets))


class OutputTests(unittest.TestCase):
    def test_pl_commentary_title(self):
        pack = pack_of(sample_pl_csv())
        self.assertTrue(to_markdown(pack, {'executive_summary': [], 'lines': []}, [], 'x').startswith('# CFO commentary: 2026-08 P&L'))

    def test_margin_points_never_show_minus_zero(self):
        self.assertEqual((_pts('-0.04'), _pts('-0.06'), _pts('0')), ('0.0 pts', '−0.1 pts', '0.0 pts'))

    def test_long_names_and_many_brands_stay_on_the_slide(self):
        long = 'Regional Sales and Customer Development'
        text = 'month,country,cost_center,cc_name,account,currency,actual,budget,forecast\n' + ''.join(
            f'2026-08,{c},{c[:2]}-{k},{c} {long},Recruitment,INR,{-k * 1000 - 5000},{300000 + k},{300000}\n'
            for c in ('India', 'Singapore', 'Australia', 'United Arab Emirates') for k in range(3))
        sales = BRANDS + ''.join(f'2026-08,India,PC-IN,Brand {i:02d},Product sales,INR,{100000 + i * 40000},100000,100000\n'
                                 f'2026-08,India,PC-IN,Brand {i:02d},Cost of goods sold,INR,50000,50000,50000\n' for i in range(15))
        data = merge_datasets(load_csv(text), load_csv(sales, 'Sales'))
        pack = lock_pack(build_pack(aggregate(data, data.dims), 'budget', D(25000), D(10), {}, (D(50000), D(5))))
        deck = Presentation(io.BytesIO(make_deck(pack, None)['pptx']))
        footer = 7.0 * 914400  # the footer line sits at 7.02 in
        for number, slide in enumerate(deck.slides, 1):
            for shape in slide.shapes:
                if number > 1 and shape.top < footer:  # the title slide has a full-page background
                    with self.subTest(slide=number, shape=shape.name):
                        self.assertLessEqual(shape.top + shape.height, footer)
        products = next(s for s in deck.slides
                        if any('sales and gross margin by brand' in getattr(x, 'text', '').lower() for x in s.shapes))
        table = next(x for x in products.shapes if x.has_table).table
        self.assertEqual(len(table.rows), 9)  # eight brands that move most, plus the header
        self.assertIn('Brand 14', [table.cell(r, 0).text for r in range(1, 9)])  # the largest gain is always shown


if __name__ == '__main__':
    unittest.main()
