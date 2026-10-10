"""Read any CSV or Excel file, map its columns to roles, and turn it into the engine's standard layout.

Two ways to supply data:
- one combined file with actual, budget and forecast columns, or
- separate actuals, budget and (optional) forecast files, matched line by line.
The standard layout then goes through finance.load_csv, so every validation rule still applies.
"""
import csv
import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook

from finance import CC, CC_ATTRS, SAMPLE_LINES, _column

IGNORE, EXTRA = 'ignore', 'extra'
ROLE_LABELS = {'month': 'Month', 'country': 'Country / entity', CC: 'Cost centre', 'account': 'Account or GL code',
               'currency': 'Currency', 'amount': 'Amount', 'actual': 'Actual amount', 'budget': 'Budget amount',
               'forecast': 'Forecast amount', 'cc_name': 'Cost centre name', 'function': 'Function', 'owner': 'Owner',
               'pl_line': 'P&L line', EXTRA: 'Extra dimension', IGNORE: 'Ignore'}
BASE_ROLES = ('month', 'country', CC, 'account', 'currency')
ROLES = {'combined': BASE_ROLES + ('actual', 'budget', 'forecast') + CC_ATTRS + ('pl_line', EXTRA, IGNORE),
         'pl': BASE_ROLES + ('actual', 'budget', 'forecast') + CC_ATTRS + ('pl_line', EXTRA, IGNORE),
         'actuals': BASE_ROLES + ('amount',) + CC_ATTRS + (EXTRA, IGNORE),
         'budget': BASE_ROLES + ('amount', EXTRA, IGNORE),
         'forecast': BASE_ROLES + ('amount', EXTRA, IGNORE)}
REQUIRED = {'combined': ('month', 'account', 'actual', 'budget', 'forecast'),
            'pl': ('month', 'account', 'actual', 'budget', 'forecast'),
            'actuals': ('month', 'account', 'amount'), 'budget': ('month', 'account', 'amount'),
            'forecast': ('month', 'account', 'amount')}
ALIASES = {
    'month': {'month', 'period', 'postingperiod', 'fiscalperiod', 'monthyear', 'date', 'postingdate', 'yearmonth'},
    'country': {'country', 'entity', 'company', 'legalentity', 'companycode', 'countryentity'},
    CC: {'costcenter', 'costcentre', 'cc', 'cccode', 'costcentercode', 'costcentrecode', 'profitcenter'},
    'account': {'account', 'glaccount', 'gl', 'glcode', 'accountcode', 'naturalaccount', 'ledgeraccount', 'accountname',
                'glaccountname', 'expensetype', 'costelement'},
    'currency': {'currency', 'ccy', 'curr', 'currencycode'},
    'actual': {'actual', 'actuals', 'actualamount', 'actualinr', 'act'},
    'budget': {'budget', 'budgetamount', 'budgetinr', 'bud', 'plan'},
    'forecast': {'forecast', 'forecastamount', 'forecastinr', 'fcst', 'outlook', 'latestestimate', 'le'},
    'amount': {'amount', 'amountinr', 'value', 'amt', 'inr', 'netamount', 'actual', 'actuals', 'actualamount',
               'budget', 'budgetamount', 'forecast', 'forecastamount', 'plan', 'fcst', 'outlook'},
    'cc_name': {'ccname', 'costcentername', 'costcentrename', 'ccdescription', 'costcenterdescription'},
    'function': {'function', 'department', 'dept'},
    'owner': {'owner', 'ccowner', 'budgetholder', 'budgetowner', 'manager'},
    'pl_line': {'plline', 'pnlline', 'pandlline', 'pllinename', 'plcategory', 'linetype', 'pnlcategory', 'plgroup',
                'pnlgroup', 'statementline'},
    EXTRA: {'project', 'region', 'businessunit', 'bu', 'product', 'productline', 'channel', 'segment', 'brand',
            'location', 'site', 'division', 'program', 'programme'},
}
MONTHS = {m: i for i, m in enumerate(('jan', 'feb', 'mar', 'apr', 'may', 'jun', 'jul', 'aug', 'sep', 'oct', 'nov', 'dec'), 1)}


def _key(header):
    return re.sub(r'[^a-z0-9]', '', str(header).lower())


def read_table(data, name):
    """(headers, rows) from a CSV or the first sheet of an Excel workbook. Blank rows are skipped."""
    if name.lower().endswith(('.xlsx', '.xlsm')):
        try:
            sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        except Exception:
            raise ValueError(f'{name}: could not open the Excel file.') from None
        rows = [['' if v is None else v for v in r] for r in sheet.iter_rows(values_only=True)]
    else:
        rows = list(csv.reader(io.StringIO(data.decode('utf-8-sig'))))
    rows = [r for r in rows if any(str(v).strip() for v in r)]
    if not rows:
        raise ValueError(f'{name}: the file is empty.')
    headers = [str(h).strip() or f'Column {i}' for i, h in enumerate(rows[0], 1)]
    return headers, rows[1:]


def table_csv(data, name):
    """A CSV or Excel upload as CSV text, for files read straight by the engine (for example the GL mapping).
    Excel numbers come back as written: GL code 510100, not 510100.0."""
    headers, rows = read_table(data, name)
    out = io.StringIO(newline='')
    writer = csv.writer(out)
    writer.writerow(headers)
    writer.writerows([_text(v) for v in row] for row in rows)
    return out.getvalue()


def guess_roles(headers, kind):
    roles, used = [], set()
    for header in headers:
        key, role = _key(header), IGNORE
        for candidate in ROLES[kind]:
            if candidate not in (EXTRA, IGNORE) and key in ALIASES.get(candidate, ()) and candidate not in used:
                role = candidate
                break
        else:
            if key in ALIASES[EXTRA]:
                role = EXTRA
        if role not in (EXTRA, IGNORE):
            used.add(role)
        roles.append(role)
    return roles


def check_roles(roles, kind, country=''):
    problems = [f'Choose a column for {ROLE_LABELS[r]}.' for r in REQUIRED[kind] if r not in roles]
    if 'country' not in roles and not country.strip():
        problems.append('Choose a Country / entity column, or type the entity name for all rows.')
    for role in set(roles) - {EXTRA, IGNORE}:
        if roles.count(role) > 1:
            problems.append(f'{ROLE_LABELS[role]} is chosen for more than one column.')
    if kind == 'combined' and any(r in roles for r in CC_ATTRS) and CC not in roles:
        problems.append('Cost centre name, function or owner need a Cost centre column.')
    return problems


def normalize_month(value):
    """YYYY-MM from dates, 2026-08, 2026-08-31, 08/2026, Aug-26, August 2026. Day-first dates are refused as ambiguous."""
    if isinstance(value, (datetime, date)):
        return f'{value.year:04d}-{value.month:02d}'
    text = str(value).strip()
    for pattern, order in ((r'(\d{4})[-/.](\d{1,2})(?:[-/.]\d{1,2})?(?:[ T].*)?', 'ym'), (r'(\d{1,2})[-/.](\d{4})', 'my'),
                           (r'([A-Za-z]{3,9})[-\s/.,\']*(\d{4}|\d{2})', 'ny'), (r'(\d{4})[-\s/.]*([A-Za-z]{3,9})', 'yn')):
        match = re.fullmatch(pattern, text)
        if not match:
            continue
        a, b = match.groups()
        if order == 'ym':
            year, month = int(a), int(b)
        elif order == 'my':
            month, year = int(a), int(b)
        else:
            name, year = (a, b) if order == 'ny' else (b, a)
            month = MONTHS.get(name[:3].lower())
            year = int(year) + (2000 if len(year) == 2 else 0)
        if month and 1 <= month <= 12:
            return f'{year:04d}-{month:02d}'
    raise ValueError(f'"{text}" is not a recognised month. Use YYYY-MM, a date, or a month name such as Aug-2026.')


def parse_amount(value, allow_negative=False):
    if isinstance(value, bool):
        raise InvalidOperation
    if isinstance(value, (int, float, Decimal)):
        amount = Decimal(str(value))
        if isinstance(value, float):  # Excel stores doubles; drop binary noise beyond cents
            rounded = amount.quantize(Decimal('.01'))
            amount = rounded if abs(amount - rounded) < Decimal('0.000001') else amount
    else:
        text = re.sub(r'[\s,]', '', str(value))
        negative = text.startswith('(') and text.endswith(')')
        amount = Decimal(text.strip('()'))
        amount = -amount if negative else amount
    if not amount.is_finite() or amount != amount.quantize(Decimal('.01')) or (amount < 0 and not allow_negative):
        raise InvalidOperation
    return amount


def _text(value):
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _cells(headers, row, roles, name, line):
    if len(row) != len(headers):
        raise ValueError(f'{name} row {line}: missing value or extra column. No values were assumed.')
    return {(h if r == EXTRA else r): v for h, r, v in zip(headers, roles, row) if r != IGNORE}


def _dims(headers, roles):
    return ([CC] if CC in roles else []) + [_column(h) for h, r in zip(headers, roles) if r == EXTRA]


def _write(rows, dims, attrs):
    out = io.StringIO(newline='')
    writer = csv.writer(out)
    writer.writerow(('month', 'country', *dims, 'account', 'currency', 'actual', 'budget', 'forecast', *attrs))
    writer.writerows(rows)
    return out.getvalue()


def combined_csv(headers, rows, roles, name, country='', currency='INR', allow_negative=True):
    """One file with actual, budget and forecast columns, rewritten in the standard layout.
    Negative amounts (credits, reversals) are kept; the review flags lines that net below zero."""
    dims, attrs = _dims(headers, roles), [a for a in CC_ATTRS + ('pl_line',) if a in roles]
    extra_cols = {h: _column(h) for h, r in zip(headers, roles) if r == EXTRA}
    out = []
    for line, row in enumerate(rows, 2):
        cells = _cells(headers, row, roles, name, line)
        values = {extra_cols.get(k, k): v for k, v in cells.items()}
        try:
            month = normalize_month(values['month']) if _text(values['month']) else ''
        except ValueError as error:
            raise ValueError(f'{name} row {line}: {error}') from None
        amounts = []
        for field in ('actual', 'budget', 'forecast'):
            raw = values[field]
            try:
                amounts.append(parse_amount(raw, allow_negative) if _text(raw) else '')
            except InvalidOperation:
                kind = 'a' if allow_negative else 'a nonnegative'
                raise ValueError(f'{name} row {line}: {field} must be {kind} number with at most 2 decimals.') from None
        out.append((month, _text(values['country']) if 'country' in values else country.strip(),
                    *(_text(values[d]) for d in dims), _text(values['account']),
                    _text(values['currency']) if 'currency' in values else currency, *amounts,
                    *(_text(values[a]) for a in attrs)))
    return _write(out, dims, attrs)


def _single(kind, headers, rows, roles, name, country, currency):
    """Sum one file's rows into lines keyed by month/country/dims/account/currency."""
    dims = _dims(headers, roles)
    extra_cols = {h: _column(h) for h, r in zip(headers, roles) if r == EXTRA}
    lines, attrs = {}, {}
    for line, row in enumerate(rows, 2):
        values = {extra_cols.get(k, k): v for k, v in _cells(headers, row, roles, name, line).items()}
        needed = ['month', 'account', 'amount'] + dims + [r for r in ('country', 'currency') if r in values]
        if any(not _text(values[k]) for k in needed):
            raise ValueError(f'{name} row {line}: missing value. No values were assumed.')
        try:
            month = normalize_month(values['month'])
            amount = parse_amount(values['amount'], allow_negative=True)
        except ValueError as error:
            raise ValueError(f'{name} row {line}: {error}') from None
        except InvalidOperation:
            raise ValueError(f'{name} row {line}: amount must be a number with at most 2 decimals.') from None
        key = (month, _text(values['country']) if 'country' in values else country.strip(),
               tuple((d, _text(values[d])) for d in dims), _text(values['account']),
               _text(values['currency']) if 'currency' in values else currency)
        lines[key] = lines.get(key, Decimal(0)) + amount
        if kind == 'actuals' and CC in dims:
            info = {a: _text(values[a]) for a in CC_ATTRS if a in values and _text(values[a])}
            if attrs.setdefault(_text(values[CC]), info) != info:
                raise ValueError(f'{name} row {line}: cost centre {_text(values[CC])} has different details than an earlier row.')
    return lines, dims, attrs


def separate_csv(files, country='', currency='INR', resolve=None):
    """files: {'actuals': (headers, rows, roles, name), 'budget': ..., optional 'forecast': ...}.

    Postings are netted per line; a line may net below zero (a credit note larger than the spend).
    resolve: GL code -> reporting account, when a GL mapping is used. Lines then stay at GL code level,
    but a line counts as budgeted when its account has a budget in any GL code,
    so a budget kept at account level matches actuals posted to GL codes.
    Returns (standard CSV text, cost centre details from the actuals, reconciliation report).
    """
    parsed = {kind: _single(kind, *spec, country, currency) for kind, spec in files.items()}
    actual_lines, actual_dims, attrs = parsed['actuals']
    months = sorted({k[0] for k in actual_lines})
    if len(months) != 1:
        raise ValueError('The actuals file must hold one month; it has ' + ', '.join(months) + '.')
    month = months[0]
    common = [d for d in actual_dims if all(d in parsed[k][1] for k in parsed)]
    dropped = [d for d in actual_dims if d not in common] + sorted({d for k in parsed for d in parsed[k][1]} - set(actual_dims))
    report = {'month': month, 'dims': common, 'dropped_dims': dropped, 'files': [], 'forecast': 'forecast' in parsed}
    merged = {}
    for kind, (lines, _, _) in parsed.items():
        used = {k: v for k, v in lines.items() if k[0] == month}
        if not used:
            raise ValueError(f"The {kind} file has no rows for {month}.")
        report['files'].append({'kind': kind, 'name': files[kind][3], 'rows': len(files[kind][1]),
                                'lines': len(used), 'other_months': len({k[0] for k in lines}) - 1,
                                'total': sum(used.values(), Decimal(0))})
        for (m, ctry, dims, account, ccy), amount in used.items():
            key = (ctry, tuple(v for d, v in dims if d in common), account, ccy)
            slot = merged.setdefault(key, {})
            slot[kind] = slot.get(kind, Decimal(0)) + amount
    rows, found = [], {}
    for (ctry, dims, account, ccy), slot in merged.items():
        found.setdefault(' | '.join((ctry, *dims, resolve(account) if resolve else account)), set()).update(slot)
        actual, budget = slot.get('actuals', Decimal(0)), slot.get('budget', Decimal(0))
        forecast = slot.get('forecast', Decimal(0)) if 'forecast' in parsed else budget
        rows.append((month, ctry, *dims, account, ccy, actual, budget, forecast))
    report.update(unbudgeted=[l for l, s in found.items() if 'budget' not in s],
                  unspent=[l for l, s in found.items() if 'actuals' not in s])
    attrs = {code: info for code, info in attrs.items() if info} if CC in common else {}
    return _write(rows, common, []), attrs, report


# Sample files in the shape real systems export. Together they reproduce the combined demo exactly.
def sample_actuals_csv():
    """GL-style export: several postings per line, one credit note, month as Aug-2026, no currency column."""
    out = io.StringIO(newline='')
    writer = csv.writer(out)
    writer.writerow(('Posting Period', 'Entity', 'Cost Center', 'GL Account', 'Description', 'Amount (INR)'))
    for country, cc, account, budget, change in SAMPLE_LINES:
        actual = budget + change
        if cc == 'IN-SAL' and account == 'Travel':
            writer.writerow(('Aug-2026', country, cc, account, 'Travel agency invoice', f'{actual + 20000:,}'))
            writer.writerow(('Aug-2026', country, cc, account, 'Credit note: cancelled trip', f'({20000:,})'))
            continue
        first = actual * 7 // 10
        writer.writerow(('Aug-2026', country, cc, account, 'Monthly posting', first))
        if actual - first:
            writer.writerow(('Aug-2026', country, cc, account, 'Accrual', actual - first))
    return out.getvalue()


def sample_budget_xlsx():
    """Full-year budget workbook, January to December 2026. Lines with no budget are simply absent."""
    book = Workbook()
    sheet = book.active
    sheet.title = 'Budget FY26'
    sheet.append(('Month', 'Country', 'CC', 'GL', 'Budget'))
    for month in range(1, 13):
        for country, cc, account, budget, _ in SAMPLE_LINES:
            if budget:
                sheet.append((f'2026-{month:02d}', country, cc, account, budget))
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def sample_forecast_csv():
    out = io.StringIO(newline='')
    writer = csv.writer(out)
    writer.writerow(('Period', 'Country', 'Cost Centre', 'Account', 'Forecast Amount'))
    for country, cc, account, budget, change in SAMPLE_LINES:
        writer.writerow(('2026-08-31', country, cc, account, budget + change // 2))
    return out.getvalue()