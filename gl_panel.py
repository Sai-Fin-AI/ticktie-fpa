"""GL mapping screens: the mapping workbook (the master, with any new GL codes added) and the GL code tables."""
import io

import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from finance import GL, PL_LINES, breakdown, dim_label, gl_detail, gl_lines
from pack import workbook_bytes
from style import colour_table, money_columns

XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
NEW = 'New: fill in the account'


def mapping_workbook(gl_map, present, new=()):
    """The GL mapping as an Excel file that uploads straight back.

    New codes come first with the account left blank, then the mapping in use. Status says Mapped, New, or
    No rows this month. GL codes are text cells, so leading zeros survive; account and P&L line have dropdowns.
    """
    book = Workbook()
    sheet = book.active
    sheet.title = 'GL mapping'
    sheet.append((GL, 'gl_name', 'account', 'pl_line', 'status'))
    present = set(present)
    for code in new:
        sheet.append((code, '', '', '', NEW))
    for code, entry in sorted(gl_map.items()):
        sheet.append((code, entry['gl_name'], entry['account'], entry.get('pl_line', ''),
                      'Mapped' if code in present else 'No rows this month'))
    header, fill = PatternFill('solid', fgColor='13294B'), PatternFill('solid', fgColor='FEF0C7')
    for cell in sheet[1]:
        cell.font, cell.fill = Font(bold=True, color='FFFFFF'), header
    for row in sheet.iter_rows(min_row=2):
        row[0].number_format = '@'
        if row[4].value == NEW:
            for cell in row:
                cell.fill = fill
    for column, width in zip('ABCDE', (14, 34, 26, 30, 24)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = 'A2'
    last = max(sheet.max_row, 2) + 200  # room for codes added by hand
    lists = book.create_sheet('Lists')
    accounts = sorted({e['account'] for e in gl_map.values()})
    for i, account in enumerate(accounts, 1):
        lists.cell(i, 1, account)
    lists.sheet_state = 'hidden'
    if accounts:  # a pick list of known accounts; typing a new account is still allowed
        pick = DataValidation(type='list', formula1=f'=Lists!$A$1:$A${len(accounts)}', showErrorMessage=False)
        sheet.add_data_validation(pick)
        pick.add(f'C2:C{last}')
    lines = DataValidation(type='list', formula1='"' + ','.join(PL_LINES) + '"', allow_blank=True)
    sheet.add_data_validation(lines)
    lines.add(f'D2:D{last}')
    notes = book.create_sheet('How to use')
    for text in ('One row per GL code. gl_code and account are required; gl_name and pl_line are optional.',
                 'account: the reporting account the GL code rolls up to. Several GL codes can share one account.',
                 'pl_line: where the account sits on the P&L (' + ', '.join(PL_LINES) + '). Blank: read from the '
                 'account name, otherwise Operating expenses. All GL codes of an account sit on one line.',
                 'Yellow rows are GL codes in the data that the mapping does not know yet: fill in their account.',
                 'The status column is for information and is ignored on upload. Save and upload this file in the sidebar.'):
        notes.append((text,))
    notes.column_dimensions['A'].width = 120
    return workbook_bytes(book)


def product_workbook(master, present, new=()):
    """The product master as Excel: new SKUs first (brand and category blank, highlighted), then the master in use.
    SKU codes are text cells, so leading zeros survive."""
    book = Workbook()
    sheet = book.active
    sheet.title = 'Product master'
    sheet.append(('sku', 'sku_name', 'brand', 'category', 'status'))
    present = set(present)
    for sku in new:
        sheet.append((sku, '', '', '', 'New: fill in the brand and category'))
    for sku, entry in sorted(master.items()):
        sheet.append((sku, entry.get('sku_name', ''), entry.get('brand', ''), entry.get('category', ''),
                      'In the data' if sku in present else 'No rows this month'))
    header, fill = PatternFill('solid', fgColor='13294B'), PatternFill('solid', fgColor='FEF0C7')
    for cell in sheet[1]:
        cell.font, cell.fill = Font(bold=True, color='FFFFFF'), header
    for row in sheet.iter_rows(min_row=2):
        row[0].number_format = '@'
        if str(row[4].value).startswith('New'):
            for cell in row:
                cell.fill = fill
    for column, width in zip('ABCDE', (14, 32, 18, 18, 34)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = 'A2'
    notes = book.create_sheet('How to use')
    for text in ('One row per SKU. sku is required, plus brand and/or category; sku_name is optional.',
                 'A brand sits in one category. The status column is for information and is ignored on upload.',
                 'Yellow rows are SKUs in the data that the master does not know yet: fill in their brand and category.'):
        notes.append((text,))
    notes.column_dimensions['A'].width = 110
    return workbook_bytes(book)


def margins_table(items, name, baseline):
    """Sales and gross margin per brand (or SKU, category)."""
    pct = lambda v: f'{v:.1f}%' if v is not None else 'N/A'
    frame = pd.DataFrame([{name: e['label'], 'Sales INR': float(e['sales_actual']),
                           f'Sales {baseline} INR': float(e['sales_baseline']), 'Sales variance INR': float(e['sales_variance']),
                           'Sales variance %': pct(e['sales_pct']), 'Gross margin INR': float(e['gm_actual']),
                           'Gross margin %': pct(e['gm_pct_actual']), f'GM % {baseline}': pct(e['gm_pct_baseline']),
                           'GM % change (pts)': f"{e['gm_pts']:+.1f}" if e['gm_pts'] is not None else 'N/A'}
                          for e in items])
    styled = frame.style.map(lambda v: 'color:#027a48;font-weight:600' if isinstance(v, float) and v > 0 else
                             'color:#b42318;font-weight:600' if isinstance(v, float) and v < 0 else '',
                             subset=['Sales variance INR'])
    st.dataframe(styled, width='stretch', hide_index=True, column_config=money_columns(frame))
    st.caption('Gross margin = sales minus cost of sales for that ' + name.lower() + '. Green = sales above plan. '
               'Lines without a ' + name.lower() + ' (for example tax) are not in this table.')


def product_breakdown(dataset, line, baseline, level):
    """The SKUs (or brands) inside the line being explained."""
    parts = breakdown(dataset, line, baseline, level)
    label = dim_label(level)
    with st.expander(f'{label}s in this line ({len(parts)})', expanded=True):
        frame = pd.DataFrame([{label: p[level], **({'Name': p['name']} if any(q['name'] for q in parts) else {}),
                               'Actual INR': float(p['actual']), f'{baseline.title()} INR': float(p[baseline]),
                               'Variance INR': float(p['variance']),
                               'Variance %': f"{p['pct']:.1f}%" if p['pct'] is not None else 'N/A', 'Outcome': p['direction']}
                              for p in parts])
        st.dataframe(colour_table(frame) if len(frame) else frame, width='stretch', hide_index=True,
                     column_config=money_columns(frame))


def _table(items, grain, baseline):
    frame = pd.DataFrame([{'Country': g['country'], **{dim_label(d): g.get(d, '') for d in grain}, 'Account': g['account'],
                           'GL code': g[GL], 'GL name': g['gl_name'], 'Actual INR': float(g['actual']),
                           f'{baseline.title()} INR': float(g[baseline]), 'Variance INR': float(g['variance']),
                           'Variance %': f"{g['pct']:.1f}%" if g['pct'] is not None else 'N/A', 'Outcome': g['direction']}
                          for g in items])
    st.dataframe(colour_table(frame) if len(frame) else frame, width='stretch', hide_index=True,
                 column_config=money_columns(frame))


def gl_codes_view(dataset, grain, baseline, in_view):
    """Variance detail by GL code: every GL code behind the lines, largest variance first."""
    _table([g for g in gl_lines(dataset, grain, baseline) if in_view(g)], grain, baseline)
    st.caption('Each GL code at the chosen level, largest variance first. A credit GL code (for example a recharge '
               'or recovery) can be negative.')


def gl_breakdown(dataset, line, baseline):
    """The GL codes inside the line being explained, so the owner sees which postings drive it."""
    parts = gl_detail(dataset, line, baseline)
    with st.expander(f'GL codes in this line ({len(parts)})', expanded=True):
        _table(parts, (), baseline)
