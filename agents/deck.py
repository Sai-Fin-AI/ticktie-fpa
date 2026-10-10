"""CFO deck agent: builds a PowerPoint review deck from a locked pack.

Every chart, table and number is drawn by code from the pack. The AI writes only the slide
headlines and three summary bullets; each one is tick-and-tied, and any that fail are replaced
by standard wording, so the deck always builds and never shows an untraced figure.
"""
import io
import json
import math
import re
from datetime import datetime
from decimal import Decimal

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

import fmt
from finance import MARGIN_NAMES
from agents import full_year_facts, inr, material_lines, owner, pct_text
from checker import allowed_figures, internal_codes, misattributed_figures, short_name, unsupported_figures
from llm import AIError
from pack import bridge_choices, load_pack, open_reason, pack_bridge, pack_json

# Navy dominates; teal is the accent; red/green/amber carry meaning only (overspend, underspend, open).
NAVY, NAVY_SOFT, NAVY_LINE = RGBColor.from_string('13294B'), RGBColor.from_string('1E3A63'), RGBColor.from_string('3A5680')
INK, SLATE, MUTED = RGBColor.from_string('1B2433'), RGBColor.from_string('56637A'), RGBColor.from_string('8A96A8')
MIST, RULE, WHITE, ICE = RGBColor.from_string('F1F4F8'), RGBColor.from_string('D9DFE7'), RGBColor.from_string('FFFFFF'), RGBColor.from_string('C9D6E8')
TEAL, TEAL_LIGHT = RGBColor.from_string('0F766E'), RGBColor.from_string('5EC4B6')
RED, GREEN, AMBER = RGBColor.from_string('B42318'), RGBColor.from_string('067647'), RGBColor.from_string('B54708')
RED_LIGHT, GREEN_LIGHT = RGBColor.from_string('FF8A80'), RGBColor.from_string('6EE7B7')
FONT = 'Calibri'
KEYS = ('summary', 'bridge', 'top', 'heatmap', 'offsets', 'owners', 'outlook')
TIMINGS = ('One-off', 'Recurring', 'Timing difference', 'Unknown')

SYSTEM = """You write slide headlines for a CFO variance review deck (expenses, or a full P&L from sales to net profit).
Each headline states the "so what" of one slide in at most 14 words.
Rules:
1. Use ONLY the facts given. Never add causes, events or judgements that are not in the facts.
2. Copy amounts exactly as written in the facts (for example "INR 380,000"). Do not calculate new numbers.
3. Costs above plan are unfavorable; sales and profit above plan are favorable. Use the outcome words in the facts.
4. Never write internal codes such as L01.
5. With cost centres, a country plus an account ("India Distribution") means the account total across cost centres.
   When a headline is about a single line, name its cost centre as written in the facts.
6. Each headline must be about its own slide's topic (listed below); do not reuse another slide's message.
Slides: summary = overall result; bridge = which accounts move the bridge measure from plan to actual;
top = largest line movements; heatmap = where variance is concentrated; offsets = what net figures hide;
owners = who holds open items; outlook = run-rate impact by timing.
Also write exactly 3 summary bullets for the executive summary slide."""

SCHEMA = {'type': 'object', 'additionalProperties': False, 'required': ['headlines', 'summary'], 'properties': {
    'headlines': {'type': 'object', 'additionalProperties': False, 'required': list(KEYS),
                  'properties': {k: {'type': 'string'} for k in KEYS}},
    'summary': {'type': 'array', 'minItems': 3, 'maxItems': 3, 'items': {'type': 'string'}}}}


def _d(value):
    return Decimal(value) if value is not None else Decimal(0)


def _whole(value):
    return f'INR {fmt.group(_d(value))}'


def _balance(value):
    """An actual, budget or profit, keeping the minus sign of a loss or a negative balance."""
    return f'INR {fmt.balance(_d(value))}'


def _compact(value):
    return fmt.compact(_d(value))


def _profit(pack):
    return pack.get('kind') == 'pl'


def bridge_walk(pack):
    """The bridge the deck shows: the measure chosen on screen (default EBITDA for a P&L, costs otherwise)."""
    measure = pack.get('_bridge') or ('EBITDA' if _profit(pack) else 'Costs')
    return pack_bridge(pack, measure if measure in bridge_choices(pack) else bridge_choices(pack)[0])


def account_steps(pack):
    """(account, change in the bridge measure) per account across all countries and cost centres, largest first.
    For an expense review this is each account's net variance."""
    return [(account, delta) for account, delta, _ in bridge_walk(pack)['steps']]


def _measure_word(measure):
    return measure.lower() if measure in ('Sales', 'Costs') else measure

def timing_split(pack):
    """Confirmed explained amounts by timing, plus what is still unexplained (material lines only)."""
    split = {t: Decimal(0) for t in TIMINGS}
    for _, line in material_lines(pack):
        if line['note'] and _d(line['explained']):
            split[line['note']['timing']] += _d(line['explained'])
    split['Unexplained'] = _d(pack['summary']['unexplained'])
    return split


def owner_table(pack):
    owners = {}
    for _, line in material_lines(pack):
        why = open_reason(line)
        if why:
            entry = owners.setdefault(owner(line), {'lines': 0, 'open': Decimal(0), 'reasons': {}})
            entry['lines'] += 1
            entry['open'] += _d(line['unexplained'])
            entry['reasons'][why] = entry['reasons'].get(why, 0) + 1
    return sorted(owners.items(), key=lambda kv: kv[1]['open'], reverse=True)


def explained_share(pack):
    gross = sum(abs(_d(l['variance'])) for l in pack['lines'] if l['material'])
    return (_d(pack['summary']['explained']) / gross * 100) if gross else None


def _total_words(pack):
    t = pack['totals']
    word = 'above' if _d(t['variance']) > 0 else 'below'
    return t.get('name', 'Expenses'), word


def _pl_facts(pack):
    """Subtotals and margins as ready phrases (P&L reviews only)."""
    out = {}
    for e in pack.get('pl') or []:
        if e['kind'] == 'subtotal' or (e['kind'] == 'line' and e['label'] == 'Sales'):
            v = _d(e['variance'])
            out[e['label']] = (f"{_balance(e['actual'])} vs {_balance(e['baseline'])}: {_whole(v)} "
                               f"{'above' if v > 0 else 'below'} plan ({e['direction'].lower()})" if v else f"{_balance(e['actual'])}, on plan")
        elif e['kind'] == 'margin':
            out[MARGIN_NAMES.get(e['label'], e['label'])] = f"{_d(e['actual_pct']):.1f}% vs {_d(e['baseline_pct']):.1f}% ({_pts(e['pts'])})"
    return out


def _facts(pack):
    t, s, basis = pack['totals'], pack['summary'], pack['basis']
    share = explained_share(pack)
    split = timing_split(pack)
    walk = bridge_walk(pack)
    name, word = _total_words(pack)
    facts = {
        'period': pack['period'], 'comparison': basis, 'review_total': name,
        'net': f"{name} {_whole(t['variance'])} {word} {basis} ({t['direction'].lower()})",
        'net_pct': pct_text(t) if t['pct'] is not None else 'N/A',
        'actual': _balance(t['actual']), basis: _balance(t['baseline']),
        'explained_share': f'{share:.0f}%' if share is not None else 'N/A',
        'explained_confirmed': _whole(s['explained']), 'unexplained': _whole(s['unexplained']),
        'lines_to_explain': s['material'], 'open_items': s['open_items'],
        'bridge_measure': walk['measure'],
        'accounts': [f"{a}: {'adds' if v > 0 else 'takes'} {_whole(v)} {'to' if v > 0 else 'off'} {_measure_word(walk['measure'])} "
                     f"({'favorable' if fav else 'unfavorable'})" for a, v, fav in walk['steps'][:4]],
        'largest_lines': [f"{l['label']}: {_whole(l['variance'])} {l['direction'].lower()}" for _, l in material_lines(pack)[:4]],
        'offsets': [f"{e['label']}: net {_whole(e['net'])}, gross {_whole(e['gross'])}"
                    for e in pack.get('rollups', {}).get('account', []) if 'Offsetting movements' in e['flags']][:3],
        'spread': [e['label'] for e in pack.get('rollups', {}).get('account', []) if 'Material in total only' in e['flags']],
        'owners_with_open_items': [f"{name}: {v['lines']} line(s), {_whole(v['open'])} open" for name, v in owner_table(pack)[:3]],
        'timing': {k: _whole(v) for k, v in split.items()},
        'full_year': full_year_facts(pack.get('outlook'), _whole),
    }
    if _profit(pack):
        facts['profit_and_loss'] = _pl_facts(pack)
    return facts


def _bridge_headline(pack):
    walk = bridge_walk(pack)
    if not walk['steps']:
        return 'Actual equals plan'
    account, v, _ = walk['steps'][0]
    if not _profit(pack):
        return f"{account} is the largest account movement at {_whole(v)}"
    return f"{account} {'adds' if v > 0 else 'takes'} {_whole(v)} {'to' if v > 0 else 'off'} {_measure_word(walk['measure'])}, the largest step"


def default_headlines(pack):
    t, s, basis = pack['totals'], pack['summary'], pack['basis']
    share, owners = explained_share(pack), owner_table(pack)
    offsets = [e for e in pack.get('rollups', {}).get('account', []) if 'Offsetting movements' in e['flags']]
    split = timing_split(pack)
    top = material_lines(pack)[0][1] if material_lines(pack) else None
    name, word = _total_words(pack)
    return {
        'summary': f"{name} {_whole(t['variance'])} {word} {basis}; "
                   + (f"{share:.0f}% of material movement explained" if share is not None else 'no material movements'),
        'bridge': _bridge_headline(pack),
        'top': (f"{short_name(top)} is the largest line: {_whole(top['variance'])} {top['direction'].lower()}" if top
                else 'No line meets the thresholds'),
        'heatmap': f"{s['material']} lines need explanation across {len({l['country'] for l in pack['lines']})} countries",
        'offsets': (f"{offsets[0]['label']}: net {_whole(offsets[0]['net'])} hides {_whole(offsets[0]['gross'])} gross movement"
                    if offsets else 'No material offsets hidden inside account totals'),
        'owners': (f"{owners[0][0]} holds the most open variance: {_whole(owners[0][1]['open'])}" if owners
                   else 'No open items'),
        'outlook': _outlook_headline(pack, split),
    }


def _outlook_headline(pack, split):
    outlook = pack.get('outlook')
    if not outlook or not _d(outlook['fy_effect']):
        return (f"Recurring {_whole(split['Recurring'])}, one-off {_whole(split['One-off'])}, "
                f"unexplained {_whole(split['Unexplained'])}")
    effect = _d(outlook['fy_effect'])
    if outlook.get('terms') == 'profit':
        return (f"Confirmed items {'add' if effect > 0 else 'take'} {_whole(effect)} {'to' if effect > 0 else 'off'} "
                f"{outlook['fy']} profit; {_whole(split['Unexplained'])} still unexplained")
    word = 'above' if effect > 0 else 'below'
    return (f"Confirmed items point to {_whole(effect)} {word} plan for {outlook['fy']}; "
            f"{_whole(split['Unexplained'])} still unexplained")


def default_summary(pack):
    t, s, basis = pack['totals'], pack['summary'], pack['basis']
    if _profit(pack):
        name, word = _total_words(pack)
        first = (f"{name} {_balance(t['actual'])} vs {basis} {_balance(t['baseline'])}: {_whole(t['variance'])} {word} "
                 f"({t['direction'].lower()}).")
    else:
        first = f"Actual {_balance(t['actual'])} vs {basis} {_balance(t['baseline'])}: net {_whole(t['variance'])} {t['direction'].lower()}."
    return [first,
            f"Confirmed explanations cover {_whole(s['explained'])}; {_whole(s['unexplained'])} remains unexplained.",
            f"{s['open_items']} open item(s) across {s['material']} lines that need explanation."]


def derived_records(pack):
    """Figures the deck shows beyond lines and roll-ups: bridge steps and ends, owner totals, timing split."""
    outlook = pack.get('outlook') or {}
    full_year = [('Full-year Recurring', [outlook['recurring_month'], outlook['recurring_rest'], outlook['recurring_fy'],
                                          outlook['fy_effect']]),
                 ('Full-year', [outlook['fy_effect']])] if outlook else []
    walk = bridge_walk(pack)
    return ([(account, [value]) for account, value in account_steps(pack)]
            + [(walk['measure'], [walk['start'], walk['end'], walk['end'] - walk['start']])]
            + [(name, [v['open']]) for name, v in owner_table(pack)]
            + [(timing, [value]) for timing, value in timing_split(pack).items() if timing != 'Unexplained']
            + full_year)

def _problems(text, allowed, pack):
    issues = [f'"{f}" is not in the pack' for f in unsupported_figures(text, allowed)]
    issues += misattributed_figures(text, pack, derived_records(pack))
    issues += [f'uses the internal code {c}' for c in internal_codes(text)]
    if len(text) > 130:
        issues.append('too long for a headline')
    return issues


def draft_headlines(pack, client):
    """Return (headlines, summary, issues, source). AI text that fails the checks is replaced, never shown."""
    defaults, fallback_summary = default_headlines(pack), default_summary(pack)
    if client is None:
        return defaults, fallback_summary, [], 'standard titles'
    allowed, facts = allowed_figures(pack), json.dumps(_facts(pack), ensure_ascii=False, indent=1)
    try:
        reply = client.complete('deck', pack['content_hash'], SYSTEM, facts, SCHEMA)

        def review(reply):
            found = {k: _problems(v, allowed, pack) for k, v in reply['headlines'].items()}
            found.update({f'summary {i}': _problems(b, allowed, pack) for i, b in enumerate(reply['summary'], 1)})
            return {k: v for k, v in found.items() if v}

        problems = review(reply)
        if problems:
            notes = '\n- '.join(f'{k}: {", ".join(v)}' for k, v in problems.items())
            reply = client.complete('deck-retry', pack['content_hash'], SYSTEM,
                                    facts + '\n\nYour previous headlines failed these checks. Fix them:\n- ' + notes, SCHEMA)
            problems = review(reply)
    except AIError as error:
        return defaults, fallback_summary, [f'AI headlines unavailable ({error}); standard titles used.'], 'standard titles'
    headlines = {k: (defaults[k] if k in problems else reply['headlines'][k].strip()) for k in KEYS}
    summary_failed = any(k.startswith('summary ') for k in problems)
    summary = fallback_summary if summary_failed else [b.strip() for b in reply['summary']]
    issues = [f'{k}: replaced with standard wording ({", ".join(v)})' for k, v in problems.items()]
    return headlines, summary, issues, f'AI ({client.label}), tick-and-tied'



# ---------- drawing helpers ----------

def _text(slide, x, y, w, h, text, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
          spacing=None, after=6, font=FONT):
    """`text` is a string or a list of lines; a line may be a list of (text, {size, bold, color}) runs."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap, frame.vertical_anchor = True, anchor
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        para.alignment, para.space_after = align, Pt(after)
        for part, style in (line if isinstance(line, list) else [(line, {})]):
            run = para.add_run()
            run.text = part
            run.font.size, run.font.bold = Pt(style.get('size', size)), style.get('bold', bold)
            run.font.color.rgb, run.font.name = style.get('color', color), font
            if spacing:
                run.font._element.set('spc', str(spacing))
    return box


def _rect(slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE, radius=0.06):
    item = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    item.fill.solid()
    item.fill.fore_color.rgb = fill
    item.line.fill.background()
    item.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        item.adjustments[0] = radius
    return item


def _badge(slide, x, y, size, fill, text, font_size=13, color=WHITE):
    """Filled circle with a number or symbol: the deck's repeating motif."""
    circle = _rect(slide, x, y, size, size, fill, MSO_SHAPE.OVAL)
    frame = circle.text_frame
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    para = frame.paragraphs[0]
    para.alignment = PP_ALIGN.CENTER
    run = para.add_run()
    run.text = text
    run.font.size, run.font.bold, run.font.color.rgb, run.font.name = Pt(font_size), True, color, FONT
    return circle


def _kicker(slide, x, y, w, text, color=TEAL):
    return _text(slide, x, y, w, 0.3, text.upper(), size=10.5, bold=True, color=color, spacing=120)


def _slide(deck, kicker, headline, footer):
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    _kicker(slide, 0.6, 0.45, 12, kicker)
    _text(slide, 0.6, 0.75, 12.13, 1.0, headline, size=26, bold=True, color=NAVY, after=0)
    _text(slide, 0.6, 7.02, 11.0, 0.25, footer, size=8.5, color=MUTED)
    _text(slide, 11.73, 7.02, 1.0, 0.25, str(len(deck.slides)), size=9, bold=True, color=SLATE, align=PP_ALIGN.RIGHT)
    return slide


def _border(cell, sides):
    """Cell borders: `sides` maps L/R/T/B to (colour, points); a missing side has no line."""
    tcPr = cell._tc.get_or_add_tcPr()
    tags = ('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB')
    for tag in tags:
        for old in tcPr.findall(qn(tag)):
            tcPr.remove(old)
    for i, (tag, side) in enumerate(zip(tags, 'LRTB')):
        line = OxmlElement(tag)
        if side in sides:
            colour, points = sides[side]
            line.set('w', str(int(Pt(points))))
            fill = OxmlElement('a:solidFill')
            rgb = OxmlElement('a:srgbClr')
            rgb.set('val', str(colour))
            fill.append(rgb)
            line.append(fill)
        else:
            line.append(OxmlElement('a:noFill'))
        tcPr.insert(i, line)


NUMERIC = re.compile(r'[+−-]?[\d,]+(\.\d+)?(%|k|m| L| Cr)?|N/A|–')


def _wrapped(text, width, size):
    """Estimated number of lines a cell's text wraps onto (a little generous, so tables never run long)."""
    per_line = max(1, int((width - 0.16) / (size * 0.0066)))  # average character width, measured on rendered slides
    return sum(max(1, -(-len(part) // per_line)) for part in str(text).split('\n'))


def _fit(header, rows, widths, size, row_h, room):
    """Row heights (header first) grown to fit wrapped text, and how many rows fit in `room` inches."""
    height = lambda cells, pt: max(row_h, max(_wrapped(t, w, pt) for t, w in zip(cells, widths)) * pt * 1.25 / 72 + 0.12)
    heights = [height(header, size - 2)]
    for row in rows:
        h = height(row, size)
        if sum(heights) + h > room:
            break
        heights.append(h)
    return heights, len(heights) - 1


def _table(slide, x, y, w, header, rows, widths, row_h=0.4, size=12, fills=None, colours=None, bolds=None):
    """Clean table: white rows, hairline rules, no vertical lines; numbers right-aligned.
    row_h: one height for every row, or a list (header first) from _fit."""
    heights = list(row_h) if isinstance(row_h, (list, tuple)) else [row_h] * (len(rows) + 1)
    shape = slide.shapes.add_table(len(rows) + 1, len(header), Inches(x), Inches(y), Inches(w), Inches(sum(heights)))
    table = shape.table
    table.first_row = table.horz_banding = False
    numeric = [c > 0 and bool(rows) and all(NUMERIC.fullmatch(str(r[c])) or str(r[c]) == '·' for r in rows)
               for c in range(len(header))]
    for i, width in enumerate(widths):
        table.columns[i].width = Inches(width)
    for r, values in enumerate([header] + rows):
        table.rows[r].height = Inches(heights[r])
        for c, value in enumerate(values):
            cell = table.cell(r, c)
            cell.text = str(value)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            para = cell.text_frame.paragraphs[0]
            para.alignment = (PP_ALIGN.CENTER if str(value) == '·' else PP_ALIGN.RIGHT if numeric[c] else PP_ALIGN.LEFT)
            font = para.runs[0].font if para.runs else para.font
            font.name = FONT
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = WHITE
                font.size, font.bold, font.color.rgb = Pt(size - 2), True, SLATE
                _border(cell, {'B': (NAVY, 1.5)})
            else:
                cell.fill.fore_color.rgb = (fills(r - 1, c) if fills else None) or WHITE
                font.size = Pt(size)
                font.color.rgb = (colours(r - 1, c) if colours else None) or INK
                font.bold = bool(bolds(r - 1, c)) if bolds else False
                for extra in cell.text_frame.paragraphs[1:]:  # a second line in a cell is a quiet sub-label
                    extra.alignment = para.alignment
                    for run in extra.runs:
                        run.font.size, run.font.bold, run.font.color.rgb, run.font.name = Pt(size - 2), False, SLATE, FONT
                _border(cell, {'B': (RULE, 0.75)})
    return table


def _blend(colour, share):
    share = max(0.0, min(1.0, share))
    return RGBColor(*(round(255 + (channel - 255) * share) for channel in colour))


def _status(line):
    if not line['note']:
        return 'Pending'
    return {'Confirmed by reviewer': 'Confirmed', 'Unconfirmed': 'Unconfirmed'}.get(line['note']['status'], 'Conflicting')


STATUS_COLOUR = {'Confirmed': GREEN, 'Unconfirmed': AMBER, 'Conflicting': RED, 'Pending': MUTED}


def _signed(value):
    return fmt.signed(_d(value))


def _tone(value):
    """Cost terms: + is spend above plan (red)."""
    return RED if _d(value) > 0 else GREEN if _d(value) < 0 else SLATE


def _dir_tone(direction):
    return {'Unfavorable': RED, 'Favorable': GREEN}.get(direction, SLATE)


def _effect_tone(value, pack):
    """Colour for a signed effect: profit terms in a P&L review (+ helps, green), cost terms otherwise."""
    value = _d(value)
    if not value:
        return SLATE
    return (GREEN if value > 0 else RED) if _profit(pack) else (RED if value > 0 else GREEN)


def _label(point, text, colour, position, size=12):
    label = point.data_label
    label.has_text_frame = True
    label.text_frame.text = text
    run = label.text_frame.paragraphs[0].runs[0]
    run.font.size, run.font.bold, run.font.color.rgb, run.font.name = Pt(size), True, colour, FONT
    label.position = position


def _tidy(chart, value_axis=True):
    """One look for every chart: no title, quiet axes, faint gridlines, deck font."""
    chart.has_title = False
    chart.font.name, chart.font.size, chart.font.color.rgb = FONT, Pt(11), SLATE
    axis = chart.value_axis
    axis.visible = value_axis
    axis.has_major_gridlines = value_axis
    if value_axis:
        axis.major_gridlines.format.line.color.rgb = RULE
        axis.format.line.fill.background()
        axis.tick_labels.font.size, axis.tick_labels.font.color.rgb = Pt(10), MUTED
    chart.category_axis.tick_labels.font.size = Pt(11)
    chart.category_axis.tick_labels.font.color.rgb = INK
    chart.category_axis.format.line.color.rgb = RULE
    chart.category_axis.has_major_gridlines = False


def _colour_series(chart, colours):
    for series, colour in zip(chart.series, colours):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = colour
        series.format.line.fill.background()


def _reason(line, limit=110):
    """The owner's explanation, shortened to fit a table cell. Pending lines say so."""
    if not line['note']:
        return 'Pending from owner'
    text = ' '.join(line['note']['reason'].split())
    return text if len(text) <= limit else text[:limit].rsplit(' ', 1)[0] + '…'


def _pct(line):
    return f"{abs(_d(line['pct'])):.1f}%" if line['pct'] is not None else 'N/A'


def _plain(value):
    return fmt.group(_d(value))


def _scaled(values):
    """Chart values in thousands (or lakhs in Indian style), so labels read 280k or 2.80 L."""
    divisor = fmt.chart_unit()[0]
    return [float(Decimal(v) / divisor) for v in values]


def _unit_format():
    return '0.00" L"' if fmt.indian() else '#,##0"k"'


def _bar(slide, x, y, w, h, parts):
    """Horizontal 100% bar from (value, colour) parts; returns the drawn shares."""
    total = sum(v for v, _ in parts if v > 0)
    left, shares = x, []
    for value, colour in parts:
        share = float(value / total) if total and value > 0 else 0.0
        if share:
            _rect(slide, left, y, w * share, h, colour)
        left += w * share
        shares.append(share)
    if not total:
        _rect(slide, x, y, w, h, RULE)
    return shares


# ---------- slides ----------

def _title_slide(deck, pack, source):
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    t, s = pack['totals'], pack['summary']
    _rect(slide, 0, 0, 13.333, 7.5, NAVY)
    _kicker(slide, 0.8, 1.55, 7.5, 'Month-end close · CFO review', color=TEAL_LIGHT)
    month = datetime.strptime(pack['period'], '%Y-%m').strftime('%B %Y')
    title = 'P&L variance review' if _profit(pack) else 'Expense variance review'
    _text(slide, 0.8, 1.95, 7.6, 2.0, [title, [(month, {'color': TEAL_LIGHT, 'size': 32})]],
          size=40, bold=True, color=WHITE, after=0)
    level = ' × '.join(['Country'] + [d.replace('_', ' ').title() for d in pack['grain']] + ['Account'])
    _text(slide, 0.8, 4.0, 7.6, 0.9, [f"Actual vs {pack['basis']} · {pack['scope']}", f"{level} · {pack['currency']}"],
          size=16, color=ICE, after=2)
    # headline number panel
    _rect(slide, 8.85, 1.55, 3.88, 3.35, NAVY_SOFT, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
    name = 'Net profit' if _profit(pack) else 'Net variance'
    _kicker(slide, 9.2, 1.85, 3.3, f"{name} vs {pack['basis']}", color=TEAL_LIGHT)
    _text(slide, 9.2, 2.2, 3.4, 0.8, _signed(t['variance']), size=38, bold=True,
          color={'Unfavorable': RED_LIGHT, 'Favorable': GREEN_LIGHT}.get(t['direction'], ICE), after=0)
    _text(slide, 9.2, 2.95, 3.4, 0.35, f"INR · {t['direction']} · {_pct(t)}", size=13, color=ICE)
    share = explained_share(pack)
    _text(slide, 9.2, 3.55, 3.4, 0.3, 'EXPLAINED (CONFIRMED)', size=10, bold=True, color=ICE, spacing=100)
    _bar(slide, 9.2, 3.9, 3.2, 0.14, [(Decimal(str(share or 0)), TEAL_LIGHT), (Decimal(100) - Decimal(str(share or 0)), NAVY_LINE)])
    _text(slide, 9.2, 4.15, 3.4, 0.35, f"{share:.0f}% · {s['open_items']} open items" if share is not None else 'No material movements',
          size=12, color=ICE)
    _text(slide, 0.8, 6.15, 11.9, 0.8,
          [f"Locked review pack {pack['content_hash'][:12]} · locked {pack.get('locked_at', 'not locked')} UTC",
           f'Figures and charts are drawn from the pack by code. Headlines: {source}. Explanations are reviewer-supplied.'],
          size=10.5, color=ICE, after=3)


def _pts(value):
    value = round(_d(value), 1)  # -0.04 shows as 0.0, not −0.0
    return f"{'+' if value > 0 else '−' if value < 0 else ''}{abs(value):.1f} pts"


def _subtotal(pack, label):
    return next((e for e in pack.get('pl') or [] if e['label'] == label), None)


def _summary_slide(deck, pack, headline, bullets, footer):
    slide = _slide(deck, 'Executive summary', headline, footer)
    t, s = pack['totals'], pack['summary']
    share = explained_share(pack)
    if _profit(pack):
        tiles = []
        for label in ('Sales', 'EBITDA', 'Net profit'):
            e, margin = _subtotal(pack, label), _subtotal(pack, f'{label} %')
            sub = f"{_signed(e['impact'])} vs {pack['basis']}" + (f" · {_d(margin['actual_pct']):.1f}% margin" if margin else '')
            tiles.append((label, ('−' if _d(e['actual']) < 0 else '') + _plain(e['actual']), sub, _dir_tone(e['direction'])))
    else:
        tiles = [('Actual', fmt.balance(_d(t['actual'])), 'INR', INK), (pack['basis'].title(), fmt.balance(_d(t['baseline'])), 'INR', INK),
                 ('Net variance', _signed(t['variance']), f"INR · {t['direction']} · {_pct(t)}", _dir_tone(t['direction']))]
    tiles += [('Explained', f'{share:.0f}%' if share is not None else '–', 'confirmed share of movement', TEAL),
              ('Open items', str(s['open_items']), f"across {s['material']} lines to explain", AMBER if s['open_items'] else GREEN)]
    width, gap = 2.226, 0.25
    for i, (label, value, sub, colour) in enumerate(tiles):
        x = 0.6 + i * (width + gap)
        _rect(slide, x, 1.95, width, 1.55, MIST, MSO_SHAPE.ROUNDED_RECTANGLE)
        _text(slide, x + 0.22, 2.12, width - 0.4, 0.3, label.upper(), size=10, bold=True, color=SLATE, spacing=100)
        _text(slide, x + 0.22, 2.42, width - 0.3, 0.6, value, size=26 if len(value) < 11 else 22, bold=True, color=colour, after=0)
        _text(slide, x + 0.22, 3.05, width - 0.35, 0.35, sub, size=10.5, color=SLATE)
    _kicker(slide, 0.6, 3.95, 7, 'Key messages')
    for i, bullet in enumerate(bullets):
        y = 4.35 + i * 0.82
        _badge(slide, 0.6, y, 0.42, NAVY, str(i + 1))
        _text(slide, 1.25, y + 0.04, 6.75, 0.75, bullet, size=15, color=INK, anchor=MSO_ANCHOR.TOP)
    # coverage panel
    _rect(slide, 8.45, 3.95, 4.28, 2.85, MIST, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04)
    _kicker(slide, 8.7, 4.15, 3.9, 'Coverage of material movement')
    explained, unexplained = _d(s['explained']), _d(s['unexplained'])
    _bar(slide, 8.7, 4.6, 3.78, 0.3, [(explained, TEAL), (unexplained, AMBER)])
    _text(slide, 8.7, 5.1, 3.8, 1.6, [
        [('■  ', {'color': TEAL, 'size': 14}), ('Confirmed  ', {'bold': True}), (_whole(explained), {})],
        [('■  ', {'color': AMBER, 'size': 14}), ('Unexplained  ', {'bold': True}), (_whole(unexplained), {})],
        [(f"{s['open_items']} open item(s) to close before the review", {'color': SLATE, 'size': 11})]], size=13, after=8)


def products_headline(pack):
    """Engine-written: the product with the largest sales gain and the one with the largest loss."""
    p = pack['products']
    rows = sorted(p['rows'], key=lambda e: _d(e['sales_variance']))
    low, high = rows[0], rows[-1]
    parts = []
    if _d(high['sales_variance']) > 0:
        parts.append(f"{high['label']} adds INR {fmt.group(_d(high['sales_variance']))} of sales")
    if _d(low['sales_variance']) < 0:
        parts.append(f"{low['label']} loses INR {fmt.group(_d(low['sales_variance']))}")
    return ('; '.join(parts) or 'Sales on plan for every ' + p['level']) + '. Mix moves inside the totals.'


def _products_slide(deck, pack, footer):
    """Sales and gross margin per brand (or SKU, category): where mix moves inside the totals."""
    p = pack['products']
    name = {'sku': 'SKU', 'brand': 'Brand', 'category': 'Category'}.get(p['level'], p['level'].title())
    slide = _slide(deck, f'Sales and gross margin by {name.lower()}', products_headline(pack), footer)
    # Eight rows fit the slide. With more products, keep the eight that move most (sales or margin), so the
    # products named in the headline are always in the table; the full list is in the Excel pack.
    moves = lambda e: abs(_d(e['sales_variance'])) + abs(_d(e['gm_variance']))
    keep = {id(e) for e in sorted(p['rows'], key=moves, reverse=True)[:8]}
    items = [e for e in p['rows'] if id(e) in keep]
    more = len(p['rows']) - len(items)
    pct = lambda v: f'{abs(_d(v)):.1f}%' if v is not None else 'N/A'
    rows = [[e['label'], fmt.balance(_d(e['sales_actual'])), _signed(e['sales_variance']), pct(e['sales_pct']),
             fmt.balance(_d(e['gm_actual'])), f"{_d(e['gm_pct_actual']):.1f}%" if e['gm_pct_actual'] is not None else 'N/A',
             _pts(e['gm_pts']) if e['gm_pts'] is not None else 'N/A'] for e in items]
    basis = pack['basis'].upper()
    tone = lambda v: GREEN if _d(v) > 0 else RED if _d(v) < 0 else SLATE
    _table(slide, 0.6, 1.95, 12.13, [name.upper(), 'SALES INR', f'VS {basis}', '%', 'GROSS MARGIN INR', 'GM %', f'GM % VS {basis}'],
           rows, [2.8, 1.8, 1.6, 1.0, 2.0, 1.2, 1.73], row_h=0.46, size=12,
           colours=lambda r, c: (tone(items[r]['sales_variance']) if c in (2, 3) else
                                 tone(items[r]['gm_pts']) if c == 6 and items[r]['gm_pts'] is not None else None),
           bolds=lambda r, c: c in (0, 2))
    _text(slide, 0.6, 1.95 + 0.46 * (len(rows) + 1) + 0.15, 12.13, 0.5,
          (f"The {len(items)} { {'Category': 'categories', 'SKU': 'SKUs'}.get(name, name.lower() + 's')} that move most; {more} more are in the Excel pack (Margins by {name.lower()} sheet). " if more else '')
          + f'Gross margin = sales minus cost of sales for that {name.lower()}. Sales: green = above plan. GM % change in '
          'percentage points. Lines without a product (corporate costs, tax) are outside this table. INR.', size=10, color=MUTED)


def _pl_slide(deck, pack, footer):
    """Sales to net profit, with subtotals in bold and margins beside the table."""
    rows = pack['pl']
    net = _subtotal(pack, 'Net profit')
    ebitda = _subtotal(pack, 'EBITDA')
    headline = (f"Net profit {_whole(net['variance'])} {'above' if _d(net['variance']) > 0 else 'below'} {pack['basis']}"
                + (f"; EBITDA {_whole(ebitda['variance'])} {'above' if _d(ebitda['variance']) > 0 else 'below'}" if ebitda else ''))
    slide = _slide(deck, 'Profit and loss', headline, footer)
    table_rows, kinds = [], []
    for e in rows:
        if e['kind'] == 'margin':
            continue
        table_rows.append([('    ' if e['kind'] == 'account' else '') + e['label'],
                           ('−' if _d(e['actual']) < 0 else '') + _plain(e['actual']),
                           ('−' if _d(e['baseline']) < 0 else '') + _plain(e['baseline']),
                           _signed(e['impact']) if _d(e['impact']) else '0', _pct(e)])
        kinds.append(e)
    row_h = min(0.3, 4.95 / (len(table_rows) + 1))
    _table(slide, 0.6, 1.85, 8.1, ['P&L LINE', 'ACTUAL INR', pack['basis'].upper() + ' INR', 'PROFIT IMPACT', '%'], table_rows,
           [3.3, 1.35, 1.35, 1.3, 0.8], row_h=row_h, size=10.5,
           fills=lambda r, c: MIST if kinds[r]['kind'] == 'subtotal' else None,
           colours=lambda r, c: (_effect_tone(kinds[r]['impact'], pack) if c == 3 else SLATE if kinds[r]['kind'] == 'account' else None),
           bolds=lambda r, c: kinds[r]['kind'] in ('subtotal', 'line') or c == 3)
    _kicker(slide, 9.1, 1.85, 3.6, 'Margins')
    y = 2.25
    for e in rows:
        if e['kind'] != 'margin':
            continue
        _rect(slide, 9.1, y, 3.63, 1.1, MIST, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
        _text(slide, 9.3, y + 0.12, 3.3, 0.3, MARGIN_NAMES.get(e['label'], e['label']).upper(), size=10, bold=True, color=SLATE, spacing=100)
        _text(slide, 9.3, y + 0.42, 1.6, 0.5, f"{_d(e['actual_pct']):.1f}%", size=22, bold=True, color=INK, after=0)
        _text(slide, 10.9, y + 0.5, 1.7, 0.4, f"{_pts(e['pts'])} vs {_d(e['baseline_pct']):.1f}%", size=11, bold=True,
              color=_dir_tone(e['direction']), align=PP_ALIGN.RIGHT)
        y += 1.25
    _text(slide, 9.1, y + 0.05, 3.63, 0.6, 'Profit impact: + helps profit (sales above plan, costs below), − hurts it.',
          size=9.5, color=MUTED)


def _bridge_slide(deck, pack, headline, footer):
    walk = bridge_walk(pack)
    measure, better = walk['measure'], walk['higher_is_better']
    kicker = (f"{measure} bridge: {pack['basis']} to actual" if _profit(pack) else f"{pack['basis'].title()} to actual bridge")
    slide = _slide(deck, kicker, headline, footer)
    steps, plan, actual = account_steps(pack), _d(walk['start']), _d(walk['end'])
    up_colour, down_colour = (GREEN, RED) if better else (RED, GREEN)
    big, big_suffix, big_word = fmt.big_unit()
    cats, base, total, up, down = [pack['basis'].title()], [0.0], [float(plan)], [0.0], [0.0]
    running, levels = plan, [plan]
    for account, delta in steps:
        cats.append(account)
        if delta >= 0:
            base.append(float(running)); up.append(float(delta)); down.append(0.0)
            running += delta
        else:
            running += delta
            base.append(float(running)); down.append(float(-delta)); up.append(0.0)
        total.append(0.0)
        levels.append(running)
    cats.append('Actual'); base.append(0.0); total.append(float(actual)); up.append(0.0); down.append(0.0)
    data = CategoryChartData()
    data.categories = cats
    for name, values in (('Base', base), ('Total', total), ('Increase', up), ('Decrease', down)):
        data.add_series(name, values)
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, Inches(0.5), Inches(1.95), Inches(8.3), Inches(4.75), data).chart
    chart.has_legend = False
    plot = chart.plots[0]
    plot.gap_width, plot.overlap = 55, 100
    series_list = list(chart.series)
    base_series = series_list[0]
    base_series.format.fill.background()
    base_series.format.line.fill.background()
    for series, colour in zip(series_list[1:], (NAVY, up_colour, down_colour)):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = colour
        series.format.line.fill.background()
    for i in range(len(cats)):
        if total[i]:  # start and end bars: amount inside, white
            _label(series_list[1].points[i], f'{Decimal(str(total[i])) / big:.2f}{big_suffix}', WHITE, XL_LABEL_POSITION.INSIDE_END)
        elif up[i] or down[i]:  # steps: amount just below the bar, on the invisible base
            _label(base_series.points[i], _compact(up[i] or -down[i]), up_colour if up[i] else down_colour,
                   XL_LABEL_POSITION.INSIDE_END)
    low, high = min(levels + [actual]), max(levels + [actual])
    span = max(high - low, Decimal(1))
    step = Decimal(10) ** int(math.floor(math.log10(float(span))))
    axis = chart.value_axis
    axis.minimum_scale = float(max(Decimal(0), ((low - span * Decimal('0.6')) / step).to_integral_value(rounding='ROUND_FLOOR') * step))
    axis.maximum_scale = float(((high + span * Decimal('0.15')) / step).to_integral_value(rounding='ROUND_CEILING') * step)
    _tidy(chart, value_axis=False)
    start = f'{Decimal(str(axis.minimum_scale)) / big:.2f}{big_suffix}'
    terms = (' Steps are each account\'s effect on ' + measure + ' (+ helps).') if _profit(pack) and measure not in ('Sales', 'Costs') else ''
    _text(slide, 0.6, 6.68, 8.2, 0.25, f'{big_word}. Bars start at {start}, not zero, so movements are visible.{terms}',
          size=9, color=MUTED)
    # what moved
    _rect(slide, 9.1, 1.95, 3.63, 4.9, MIST, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04)
    _kicker(slide, 9.35, 2.15, 3.2, 'What moved')
    shown = walk['steps'][:5]
    for i, (account, value, favorable) in enumerate(shown):
        y = 2.6 + i * 0.68
        _text(slide, 9.35, y, 1.95, 0.6, [account, [('Favorable' if favorable else 'Unfavorable', {'size': 10, 'color': SLATE})]],
              size=13, bold=True, color=INK, after=0)
        _text(slide, 11.1, y, 1.4, 0.4, _signed(value), size=15, bold=True, color=GREEN if favorable else RED, align=PP_ALIGN.RIGHT)
    y = 2.6 + len(shown) * 0.68 + 0.1
    net = actual - plan
    _text(slide, 9.35, y, 1.95, 0.4, f'Net change' if _profit(pack) else 'Net variance', size=13, bold=True, color=NAVY)
    _text(slide, 11.1, y, 1.4, 0.4, _signed(net), size=15, bold=True,
          color=SLATE if not net else GREEN if (net > 0) == better else RED, align=PP_ALIGN.RIGHT)
    if len(steps) > len(shown):
        _text(slide, 9.35, y + 0.45, 3.2, 0.3, f'{len(steps) - len(shown)} smaller account(s) in the chart only', size=10, color=SLATE)

def _top_slide(deck, pack, headline, footer):
    slide = _slide(deck, 'Largest movements', headline, footer)
    lines = material_lines(pack)[:7]
    if not lines:
        _text(slide, 0.6, 2.4, 12, 0.5, 'No line meets the materiality thresholds.', size=16, color=SLATE)
        return

    def where(line):
        parts = [line['country'], *line.get('dims', {}).values(), line['attrs'].get('owner')]
        return ' · '.join(p for p in parts if p)

    rows = [[f"{short_name(l)}\n{where(l)}" + (' · ⚠ negative balance' if l.get('negative') else ''), _signed(l['variance']), _pct(l), '● ' + _status(l), _plain(l['unexplained']),
             _reason(l)] for _, l in lines]
    header, widths = ['LINE', 'VARIANCE INR', '%', 'STATUS', 'UNEXPLAINED INR', 'REASON GIVEN'], [3.15, 1.35, 0.8, 1.4, 1.35, 4.08]
    heights, fits = _fit(header, rows, widths, 11, 0.56, 4.75)  # long centre names wrap: show the rows that fit
    rows, lines = rows[:fits], lines[:fits]
    _table(slide, 0.6, 1.85, 12.13, header, rows, widths, row_h=heights, size=11,
           colours=lambda r, c: (_dir_tone(lines[r][1]['direction']) if c == 1 else STATUS_COLOUR[_status(lines[r][1])] if c == 3
                                 else AMBER if c == 4 and rows[r][4] != '0' else MUTED if c == 4 or (c == 5 and not lines[r][1]['note'])
                                 else None),
           bolds=lambda r, c: c in (0, 1))
    _text(slide, 0.6, 1.85 + sum(heights) + 0.12, 12.13, 0.25,
          'Reasons are as supplied by owners; the status shows whether a reviewer has confirmed them. '
          f"The {len(material_lines(pack))} lines needing explanation are all in the appendix.", size=9, color=MUTED)


def _heatmap_slide(deck, pack, headline, footer):
    slide = _slide(deck, 'Where the variance sits', headline, footer)
    moving = {l['account'] for l in pack['lines'] if _d(l['variance'])}  # accounts on plan everywhere add no colour
    accounts = [a for a in dict.fromkeys(l['account'] for l in pack['lines']) if a in moving] or \
        list(dict.fromkeys(l['account'] for l in pack['lines']))
    names = list(dict.fromkeys(l['label'].rsplit(' | ', 1)[0] for l in pack['lines']))
    cells = {(l['label'].rsplit(' | ', 1)[0], l['account']): l for l in pack['lines']}
    shown = names[:10]  # ten rows leave room for the legend above the footer
    peak = max((abs(_d(l['variance'])) for l in pack['lines']), default=Decimal(1)) or Decimal(1)
    effect = lambda l: _d(l.get('impact')) if _profit(pack) else _d(l['variance'])  # Net mixes sales and costs in a P&L
    nets = {n: sum((effect(cells[(n, a)]) for a in accounts if (n, a) in cells), Decimal(0)) for n in shown}
    first_line = {n: next(l for l in pack['lines'] if l['label'].rsplit(' | ', 1)[0] == n) for n in shown}

    def row_name(n):  # a centre plus a brand would wrap onto two lines: use the codes ("India | PC-IN | Aurora")
        line = first_line[n]
        return ' | '.join([line['country'], *map(str, line['dims'].values())]) if len(line['dims']) > 1 else n

    rows = [[row_name(n)] + [_compact(cells[(n, a)]['variance']) if (n, a) in cells else '·' for a in accounts]
            + [_compact(nets[n])] for n in shown]
    wide = len(accounts) > 6  # many accounts (a P&L): shorter headers in normal case so names do not break
    first, size = (2.95, 10.5) if wide else (3.5, 11.5)
    header = ['LINE'] + [a if wide else a.upper() for a in accounts] + ['NET']
    widths = [first] + [round((12.13 - first - 1.15) / len(accounts), 2)] * len(accounts) + [1.15]
    heights, fits = _fit(header, rows, widths, size, 0.4, 4.7)  # long names or headers wrap: keep the legend on the slide
    shown, rows = shown[:fits], rows[:fits]

    def fill(r, c):
        if c == 0 or c > len(accounts):
            return None
        line = cells.get((shown[r], accounts[c - 1]))
        if not line or not _d(line['variance']):
            return WHITE
        value = _d(line['variance'])
        return _blend(RED if line['direction'] == 'Unfavorable' else GREEN, 0.1 + 0.75 * float(abs(value) / peak))

    def colour(r, c):
        if c == len(accounts) + 1:
            return _effect_tone(nets[shown[r]], pack)
        line = cells.get((shown[r], accounts[c - 1])) if 0 < c <= len(accounts) else None
        if line and abs(_d(line['variance'])) / peak > Decimal('0.5'):
            return WHITE
        return MUTED if line is None and c else None

    def bold(r, c):
        line = cells.get((shown[r], accounts[c - 1])) if 0 < c <= len(accounts) else None
        return (line and line['material']) or c == len(accounts) + 1

    _table(slide, 0.6, 1.95, 12.13, header, rows, widths, row_h=heights, size=size, fills=fill, colours=colour, bolds=bold)
    legend_y = 1.95 + sum(heights) + 0.2
    _rect(slide, 0.6, legend_y + 0.04, 0.22, 0.16, _blend(RED, 0.6))
    _rect(slide, 0.85, legend_y + 0.04, 0.22, 0.16, _blend(GREEN, 0.6))
    note = ('Unfavorable / favorable; darker = larger · bold = needs explanation · a dot means no data · INR'
            + (' · Net = profit impact (+ helps)' if _profit(pack) else ''))
    if len(names) > len(shown):
        note += f' · showing {len(shown)} of {len(names)} rows; the appendix lists every line needing explanation'
    _text(slide, 1.2, legend_y, 11.5, 0.3, note, size=10, color=SLATE)


def _offsets_slide(deck, pack, headline, footer):
    slide = _slide(deck, 'What net figures hide', headline, footer)
    entries = [e for e in pack.get('rollups', {}).get('account', []) if e['flags']][:6]
    if not entries:
        _text(slide, 0.6, 2.2, 12, 0.6, 'No account total hides offsetting lines or a spread variance at the current thresholds.',
              size=16, color=SLATE)
        return
    data = CategoryChartData()
    shown = entries[::-1]  # bar charts draw the first category at the bottom
    data.categories = [e['label'] for e in shown]
    # bar charts also draw the first series at the bottom: gross first puts net on top in each pair
    data.add_series('Gross movement', _scaled(_d(e['gross']) for e in shown))
    data.add_series('Net movement (size; signs in the cards)', _scaled(abs(_d(e['net'])) for e in shown))
    chart = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.5), Inches(1.95), Inches(7.3), Inches(4.85), data).chart
    chart.has_legend, chart.legend.position, chart.legend.include_in_layout = True, XL_LEGEND_POSITION.TOP, False
    chart.legend.font.size, chart.legend.font.name = Pt(11), FONT
    _colour_series(chart, (AMBER, NAVY))
    plot = chart.plots[0]
    plot.gap_width, plot.overlap = 60, -10
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.font.size, labels.font.bold, labels.font.name = Pt(11), True, FONT
    labels.number_format, labels.number_format_is_linked, labels.position = _unit_format(), False, XL_LABEL_POSITION.OUTSIDE_END
    _tidy(chart, value_axis=False)
    chart.value_axis.maximum_scale = max(_scaled(_d(e['gross']) for e in shown)) * 1.2 or None  # room for the labels
    y = 1.95
    for e in entries[:4]:
        offset = 'Offsetting movements' in e['flags']
        body = (f"Net {_signed(e['net'])}, gross {_plain(e['gross'])}: offsetting lines hide {_plain(e['hidden'])}." if offset
                else 'Material in total, but no single line crosses the threshold, so every owner explains their share.')
        _rect(slide, 8.1, y, 4.63, 1.13, MIST, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
        _badge(slide, 8.28, y + 0.18, 0.4, AMBER if offset else TEAL, '⇅' if offset else 'Σ', font_size=14)
        _text(slide, 8.85, y + 0.14, 3.75, 0.9, [e['label'], [(body, {'size': 11, 'bold': False, 'color': SLATE})]],
              size=13, bold=True, color=INK, after=2)
        y += 1.25


def _owners_slide(deck, pack, headline, footer):
    slide = _slide(deck, 'Open items by owner', headline, footer)
    owners = owner_table(pack)[:8]
    if not owners:
        _text(slide, 0.6, 2.2, 12, 0.6, 'Every line that needs an explanation has a confirmed one.', size=16, color=GREEN)
        return
    data = CategoryChartData()
    shown = owners[::-1]
    data.categories = [name for name, _ in shown]
    data.add_series('Unexplained INR', _scaled(v['open'] for _, v in shown))
    chart = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.5), Inches(1.95), Inches(5.6), Inches(4.85), data).chart
    chart.has_legend = False
    _colour_series(chart, (AMBER,))
    plot = chart.plots[0]
    plot.gap_width = 70
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.font.size, labels.font.bold, labels.font.name = Pt(11), True, FONT
    labels.number_format, labels.number_format_is_linked, labels.position = _unit_format(), False, XL_LABEL_POSITION.OUTSIDE_END
    _tidy(chart, value_axis=False)
    chart.value_axis.maximum_scale = max(_scaled(v['open'] for _, v in shown)) * 1.2 or None  # room for the labels
    rows = [[name, str(v['lines']), '; '.join(f'{n} × {why.lower()}' for why, n in v['reasons'].items())] for name, v in owners]
    _table(slide, 6.45, 1.95, 6.28, ['OWNER', 'LINES', 'WHY STILL OPEN'], rows, [1.75, 0.75, 3.78], row_h=0.52, size=11,
           bolds=lambda r, c: c == 0)


def _outlook_slide(deck, pack, headline, footer):
    slide = _slide(deck, 'Outlook: run-rate impact', headline, footer)
    split = timing_split(pack)
    outlook = pack.get('outlook') or {}
    # Signed (+ overspend, − saving) where the pack has it, so a recurring saving is never shown in red.
    signed = {'Recurring': outlook.get('recurring_month'), 'Timing difference': outlook.get('timing'),
              'One-off': outlook.get('one_off')}
    cards = [('↻', 'Recurring', split['Recurring'], RED, 'Likely to continue: reflect it in the next forecast'),
             ('⇄', 'Timing difference', split['Timing difference'], SLATE, 'Expected to reverse in later months'),
             ('1', 'One-off', split['One-off'], TEAL, 'No run-rate impact expected'),
             ('?', 'Unexplained', split['Unexplained'], AMBER, 'No confirmed explanation yet')]
    width, gap = 2.845, 0.25
    for i, (symbol, label, value, colour, meaning) in enumerate(cards):
        x = 0.6 + i * (width + gap)
        _rect(slide, x, 1.85, width, 1.6, MIST, MSO_SHAPE.ROUNDED_RECTANGLE)
        _badge(slide, x + 0.22, 2.0, 0.4, colour, symbol, font_size=14)
        _text(slide, x + 0.75, 2.08, width - 0.85, 0.3, label.upper(), size=10, bold=True, color=SLATE, spacing=100)
        if signed.get(label) is not None and _d(signed[label]):
            shown, tone = [(_signed(signed[label]), {}), (' INR', {'size': 12, 'bold': False, 'color': SLATE})], _effect_tone(signed[label], pack)
        else:
            shown, tone = [(_whole(value), {})], AMBER if label == 'Unexplained' else SLATE
        _text(slide, x + 0.22, 2.5, width - 0.35, 0.45, [shown], size=20, bold=True, color=tone, after=0)
        _text(slide, x + 0.22, 2.95, width - 0.4, 0.45, meaning, size=10.5, color=SLATE)
    _bar(slide, 0.6, 3.62, 12.13, 0.16, [(v, c) for _, _, v, c, _ in cards])
    note = 'Share of material movement by timing (confirmed explanations), with the unexplained remainder.'
    if split['Unknown']:
        note += f" A further {_whole(split['Unknown'])} is confirmed with unknown timing."
    _text(slide, 0.6, 3.84, 12.13, 0.3, note, size=9.5, color=MUTED)
    outlook = pack.get('outlook')
    if outlook:
        _full_year_strip(slide, outlook)
    actions = [l for _, l in material_lines(pack) if l['note'] and l['note']['action']][:4]
    _kicker(slide, 0.6, 5.55, 12, 'Actions proposed by reviewers')
    if not actions:
        _text(slide, 0.6, 5.88, 12, 0.4, 'No actions proposed yet.', size=12, color=SLATE)
    for i, line in enumerate(actions):
        x, y = 0.6 + (i % 2) * 6.2, 5.88 + (i // 2) * 0.55
        _badge(slide, x, y, 0.28, NAVY, str(i + 1), font_size=10)
        _text(slide, x + 0.4, y - 0.01, 5.65, 0.55, [[(short_name(line) + ': ', {'bold': True}), (line['note']['action'], {})]],
              size=11, color=INK, after=0)


def _full_year_strip(slide, outlook):
    """Full-year effect of confirmed explanations, with the run-rate assumption spelled out."""
    profit = outlook.get('terms') == 'profit'
    _rect(slide, 0.6, 4.2, 12.13, 1.05, MIST, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
    months = outlook['months_left']
    _kicker(slide, 0.85, 4.3, 11.6, f"Full-year view · {outlook['fy']} · {months} month{'s' if months != 1 else ''} left")
    once = _d(outlook['one_off']) + _d(outlook['unknown_timing'])  # unknown timing is counted once, never projected
    parts = [('One-off (once)' if not _d(outlook['unknown_timing']) else 'One-off or unknown timing (once)', once),
             (f'Recurring × {months + 1} months', _d(outlook['recurring_fy'])),
             ('Timing (reverses in year)', Decimal(0)),
             ('Full-year effect, confirmed', _d(outlook['fy_effect']))]
    for i, (label, value) in enumerate(parts):
        x = 0.85 + i * 3.0
        last = i == len(parts) - 1
        _text(slide, x, 4.6, 2.85, 0.25, label, size=10, bold=last, color=NAVY if last else SLATE)
        tone = (GREEN if value > 0 else RED) if profit else (RED if value > 0 else GREEN)
        _text(slide, x, 4.82, 2.85, 0.4, _signed(value) if value else '0', size=17 if last else 15, bold=True,
              color=tone if value else SLATE, after=0)
    sign_note = '+ = helps profit (favorable)' if profit else '+ = above plan (unfavorable)'
    _text(slide, 0.6, 5.27, 12.13, 0.25, 'Assumes recurring items continue at this month\'s amount; timing differences reverse '
          f'within the year. Unexplained variance is not projected. {sign_note}. INR.', size=9, color=MUTED)


def _appendix_slides(deck, pack, footer):
    lines = material_lines(pack)
    for start in range(0, max(len(lines), 1), 12):
        chunk = lines[start:start + 12]
        slide = _slide(deck, 'Appendix', f'Lines needing explanation ({start + 1}–{start + len(chunk)} of {len(lines)})', footer)
        rows = [[l['label'] + (' · ⚠ negative balance' if l.get('negative') else ''), fmt.balance(_d(l['actual'])), fmt.balance(_d(l['baseline'])), _signed(l['variance']), _pct(l), '● ' + _status(l),
                 _plain(l['unexplained'])] for _, l in chunk]
        _table(slide, 0.6, 1.75, 12.13, ['LINE', 'ACTUAL INR', pack['basis'].upper() + ' INR', 'VARIANCE INR', '%', 'STATUS',
                                         'UNEXPLAINED INR'],
               rows, [4.3, 1.4, 1.4, 1.4, 0.93, 1.4, 1.3], row_h=0.38, size=10.5,
               colours=lambda r, c, chunk=chunk: (_dir_tone(chunk[r][1]['direction']) if c == 3 else
                                                   STATUS_COLOUR[_status(chunk[r][1])] if c == 5 else
                                                   AMBER if c in (1, 2) and _d(chunk[r][1]['actual' if c == 1 else 'baseline']) < 0
                                                   else None))


def build_deck(pack, headlines, summary, source):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    footer = f"{pack['period']} · actual vs {pack['basis']} · locked pack {pack['content_hash'][:12]} · figures from the engine · headlines: {source}"
    _title_slide(deck, pack, source)
    _summary_slide(deck, pack, headlines['summary'], summary, footer)
    if _profit(pack) and pack.get('pl'):
        _pl_slide(deck, pack, footer)
    _bridge_slide(deck, pack, headlines['bridge'], footer)
    if _profit(pack) and pack.get('products'):
        _products_slide(deck, pack, footer)
    _top_slide(deck, pack, headlines['top'], footer)
    _heatmap_slide(deck, pack, headlines['heatmap'], footer)
    if pack['grain']:
        _offsets_slide(deck, pack, headlines['offsets'], footer)
    _owners_slide(deck, pack, headlines['owners'], footer)
    _outlook_slide(deck, pack, headlines['outlook'], footer)
    _appendix_slides(deck, pack, footer)
    output = io.BytesIO()
    deck.save(output)
    return output.getvalue()


def make_deck(pack, client, style=None, bridge=None):
    """Locked pack in, PowerPoint out. Never raises for AI problems: standard titles are used instead.

    bridge: the measure for the bridge slide (EBITDA, Net profit, Gross margin, Sales or Costs), as chosen on screen.
    """
    fmt.use(style)
    pack = load_pack(pack_json(pack))  # refuse a pack edited after locking
    pack = {**pack, '_bridge': bridge}  # a display choice, so it is not part of the fingerprint
    headlines, summary, issues, source = draft_headlines(pack, client)
    title = 'P&L variance review' if _profit(pack) else 'Expense variance review'
    outline = [('Title', f"{title}: {pack['period']}"), ('Executive summary', headlines['summary'])]
    if _profit(pack) and pack.get('pl'):
        outline.append(('Profit and loss', 'Sales to net profit, with subtotals and margins'))
    outline.append(('Bridge', headlines['bridge']))
    if _profit(pack) and pack.get('products'):
        outline.append(('Products', products_headline(pack)))
    outline += [('Largest movements', headlines['top']), ('Heatmap', headlines['heatmap'])]
    if pack['grain']:
        outline.append(('What net figures hide', headlines['offsets']))
    outline += [('Open items by owner', headlines['owners']), ('Outlook', headlines['outlook']), ('Appendix', 'Lines needing explanation')]
    appendix = max(1, math.ceil(len(material_lines(pack)) / 12))
    return {'pptx': build_deck(pack, headlines, summary, source), 'outline': outline, 'summary': summary,
            'slides': len(outline) - 1 + appendix,
            'issues': issues, 'source': source}
