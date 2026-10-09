"""Bulk explanation entry: an editable table of every line, and an Excel template owners fill in offline.

Both paths go through evidence.make_note, so they apply exactly the same checks as the one-line form.
A line whose explanation and reference are both left blank is treated as "no explanation".
"""
import hashlib
import io
import json
from decimal import Decimal

import pandas as pd
import streamlit as st
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from evidence import STATUSES, TIMINGS, make_note
from finance import line_label
from pack import export_notes

EDITABLE = ('Explanation', 'Reference', 'Amount explained (INR)', 'Review status', 'Timing', 'Proposed action')


def _blank(value):
    return value is None or (isinstance(value, float) and pd.isna(value)) or not str(value).strip()


def _choice(value, choices, default):
    """Match a dropdown value case-insensitively; blank means the default."""
    if _blank(value):
        return default
    found = [c for c in choices if c.lower() == str(value).strip().lower()]
    return found[0] if found else str(value).strip()


def _amount(value):
    if _blank(value):
        return '0'
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).replace(',', '').strip()


def apply_rows(records, options, baseline, notes):
    """records: [{'key': note key, Explanation..., Proposed action}]. Returns (updated notes, errors, changed count).

    Rows are applied one by one; a row with an error keeps its previous note, so one bad row never
    blocks the others.
    """
    updated, errors, changed = dict(notes), [], 0
    for record in records:
        key = tuple(record['key'])
        row = options.get(key)
        if row is None:
            errors.append(f"{record.get('Line', 'A row')}: does not match a line that needs explanation in the current data.")
            continue
        reason, reference = record.get('Explanation'), record.get('Reference')
        if _blank(reason) and _blank(reference):
            if key in updated:
                del updated[key]
                changed += 1
            continue
        try:
            note = make_note(row, baseline, '' if _blank(reason) else str(reason), '' if _blank(reference) else str(reference),
                             _amount(record.get('Amount explained (INR)')),
                             _choice(record.get('Review status'), STATUSES, STATUSES[0]),
                             _choice(record.get('Timing'), TIMINGS, TIMINGS[0]),
                             '' if _blank(record.get('Proposed action')) else str(record.get('Proposed action')))
        except ValueError as error:
            errors.append(f'{line_label(row)}: {error}')
            continue
        if updated.get(key) != note:
            updated[key] = note
            changed += 1
    return updated, errors, changed


def _records(options, notes):
    out = []
    for key, row in options.items():
        note = notes.get(key) or {}
        out.append({'key': list(key), 'Line': line_label(row), 'Owner': row.get('owner', ''),
                    'Variance (INR)': float(row['variance']),
                    'Explanation': note.get('reason', ''), 'Reference': note.get('reference', ''),
                    'Amount explained (INR)': str(note['amount']) if note else '',
                    'Review status': note.get('status', STATUSES[0]), 'Timing': note.get('timing', TIMINGS[0]),
                    'Proposed action': note.get('action', '')})
    return out


# ---------- Excel round-trip ----------

def explanation_template(options, notes, period, baseline):
    """A workbook with one row per line to explain, current notes pre-filled, dropdowns for status and timing."""
    book = Workbook()
    ws = book.active
    ws.title = 'Explanations'
    header = ['Line key (do not edit)', 'Line', 'Owner', 'Variance (INR)', *EDITABLE]
    ws.append(header)
    for record in _records(options, notes):
        ws.append([json.dumps(record['key']), record['Line'], record['Owner'], record['Variance (INR)'],
                   *(record[c] for c in EDITABLE)])
    fill, edit_fill = PatternFill('solid', fgColor='13294B'), PatternFill('solid', fgColor='EEF4FF')
    for cell in ws[1]:
        cell.font, cell.fill = Font(bold=True, color='FFFFFF'), fill
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    widths = {'A': 8, 'B': 44, 'C': 16, 'D': 14, 'E': 50, 'F': 28, 'G': 14, 'H': 24, 'I': 18, 'J': 40}
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    ws.column_dimensions['A'].hidden = True
    last = ws.max_row
    for row in ws.iter_rows(min_row=2, max_row=last):
        row[3].number_format = '#,##0;-#,##0'
        for cell in row[4:]:
            cell.fill = edit_fill
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    for column, choices in (('H', STATUSES), ('I', TIMINGS)):
        rule = DataValidation(type='list', formula1='"' + ','.join(choices) + '"', allow_blank=True)
        rule.error, rule.errorTitle = 'Pick a value from the list.', 'Not in the list'
        ws.add_data_validation(rule)
        rule.add(f'{column}2:{column}{max(last, 2)}')
    ws.freeze_panes = 'C2'
    guide = book.create_sheet('How to fill in')
    for line in (f'Explanations for {period}, actual vs {baseline}.',
                 'Fill in the blue columns only. Do not edit, add or delete rows, and do not change the hidden first column.',
                 'Explanation and Reference are both required for a line to count. Leave both blank for "no explanation yet".',
                 'Amount explained: a positive number, no larger than the variance (ignore the sign of the variance).',
                 'Review status: only "Confirmed by reviewer" counts toward explained variance.',
                 'Timing: One-off, Recurring (rolls into the full-year view) or Timing difference (expected to reverse).',
                 'Upload the saved file in the app under Business explanations → Excel round-trip.'):
        guide.append([line])
    guide.column_dimensions['A'].width = 110
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def read_template(data):
    """Rows of a filled-in template as records for apply_rows."""
    try:
        ws = load_workbook(io.BytesIO(data), data_only=True)['Explanations']
    except Exception:
        raise ValueError('This is not an explanation template from this app (no "Explanations" sheet).') from None
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h).strip() if h is not None else '' for h in rows[0]] if rows else []
    if not header or header[0] != 'Line key (do not edit)' or any(c not in header for c in EDITABLE):
        raise ValueError('The template columns were changed. Download a fresh template and copy your text into it.')
    records = []
    for number, values in enumerate(rows[1:], start=2):
        record = dict(zip(header, values))
        if _blank(record['Line key (do not edit)']):
            continue
        try:
            record['key'] = json.loads(record['Line key (do not edit)'])
        except (TypeError, json.JSONDecodeError):
            raise ValueError(f'Row {number}: the hidden line key was edited. Download a fresh template.') from None
        records.append(record)
    return records


# ---------- screens ----------

def _commit(updated, changed, errors, source):
    if errors:
        st.error(f'{len(errors)} line(s) were not saved; their previous explanation (if any) is kept:\n'
                 + '\n'.join(f'- {e}' for e in errors))
    if changed:
        st.session_state.business_notes = updated
        from ai_panel import reset_note_widgets  # late import: ai_panel imports the pack module too
        reset_note_widgets()
        st.session_state.notes_flash = f'{source}: {changed} explanation(s) saved.' + (
            f' {len(errors)} line(s) need fixing; see below.' if errors else '')
        if errors:
            st.session_state.bulk_errors = errors
        st.rerun()
    elif not errors:
        st.info(f'{source}: nothing changed.')


def notes_table(options, baseline):
    st.caption('Edit any line, then save once. Blank explanation and reference = no explanation. '
               'Only "Confirmed by reviewer" counts toward explained variance.')
    records = _records(options, st.session_state.business_notes)
    frame = pd.DataFrame(records).drop(columns=['key'])
    version = hashlib.sha256(export_notes(st.session_state.business_notes).encode()).hexdigest()[:10]
    with st.form('notes_table_form', enter_to_submit=False):
        edited = st.data_editor(
            frame, key='notes_table_' + version, hide_index=True, width='stretch', num_rows='fixed',
            disabled=['Line', 'Owner', 'Variance (INR)'],
            column_config={
                'Line': st.column_config.TextColumn(width='large'),
                'Variance (INR)': st.column_config.NumberColumn(format='localized'),
                'Explanation': st.column_config.TextColumn(width='large', max_chars=2000),
                'Reference': st.column_config.TextColumn(max_chars=300),
                'Amount explained (INR)': st.column_config.TextColumn(help='Positive number, up to the variance.'),
                'Review status': st.column_config.SelectboxColumn(options=list(STATUSES), required=True),
                'Timing': st.column_config.SelectboxColumn(options=list(TIMINGS), required=True),
                'Proposed action': st.column_config.TextColumn(max_chars=1000)})
        saved = st.form_submit_button('Save all lines', type='primary')
    if saved:
        rows = [{**r, 'key': record['key']} for r, record in zip(edited.to_dict('records'), records)]
        _commit(*apply_rows(rows, options, baseline, st.session_state.business_notes), source='Table')


def excel_roundtrip(options, baseline, period):
    st.caption('1) Download the template: one row per line, current explanations filled in, dropdowns for status and timing. '
               '2) Send it to the owners. 3) Upload the completed file here. Every row is checked like the form.')
    st.download_button('Download explanation template (Excel)',
                       explanation_template(options, st.session_state.business_notes, period, baseline),
                       f'explanations_template_{period}_{baseline}.xlsx',
                       'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    upload = st.file_uploader('Upload the completed template', type=['xlsx'], key='notes_template_upload')
    if upload is None:
        return
    digest = hashlib.sha256(upload.getvalue()).hexdigest()
    if st.session_state.get('template_import') == digest:
        st.caption('This file has been imported. Upload a new version to import again.')
        return
    st.session_state.template_import = digest
    try:
        records = read_template(upload.getvalue())
    except ValueError as error:
        st.error(str(error))
        return
    _commit(*apply_rows(records, options, baseline, st.session_state.business_notes), source='Excel template')


def show_bulk_errors():
    """Messages from the last table save or template upload, shown next to the explanation tabs."""
    if 'notes_flash' in st.session_state:
        st.success(st.session_state.pop('notes_flash'))
    for error in st.session_state.pop('bulk_errors', []):
        st.warning(error)
